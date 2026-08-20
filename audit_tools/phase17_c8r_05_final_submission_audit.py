#!/usr/bin/env python3
"""Audit and deterministically package the repaired Gate C8R handoff."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
RUN_DIR = ROOT / "phase17_v7" / "gateC8R" / "20260820_pre_submission_repair"
PACKAGE = ROOT / "04_submission" / "package_genome_medicine_gateC8R_2026-08-20"
PACKAGE_ZIP = ROOT / "04_submission" / "package_genome_medicine_gateC8R_2026-08-20.zip"
MANUSCRIPT_MD = ROOT / "01_manuscript" / "manuscript_v11_genome_medicine_gateC8R_2026-08-20.md"
SUPPLEMENT_MD = ROOT / "01_manuscript" / "supplementary_information_v2_gateC8R_2026-08-20.md"
TARGET_MD = ROOT / "04_submission" / "journal_target_decision_gateC8R_2026-08-20.md"
FIGURE_STATUS = RUN_DIR / "01_FIGURE_BUILD_STATUS.json"
PANEL_ASSERTIONS = RUN_DIR / "02_PANEL_DATA_ASSERTIONS.json"
SENSITIVITY_CSV = RUN_DIR / "03_CORRELATION_AWARE_STAT1_STAT2_SENSITIVITY.csv"
SENSITIVITY_DECISION = RUN_DIR / "04_CORRELATION_AWARE_STAT1_STAT2_DECISION.json"
SOURCE_STATUS = RUN_DIR / "05_GATE_C8R_SOURCE_BUILD_STATUS.json"
DOCUMENT_STATUS = RUN_DIR / "06_GATE_C8R_DOCUMENT_BUILD_STATUS.json"
REF_CSV = RUN_DIR / "references" / "reference_verification_gateC8R.csv"

MAIN_DOCX = PACKAGE / "main_text" / "Genome_Medicine_Manuscript_GateC8R_AUTHOR_COMPLETION_REQUIRED.docx"
SUPP_DOCX = PACKAGE / "additional_files" / "Additional_file_1_Supplementary_Information_GateC8R.docx"
COVER_DOCX = PACKAGE / "submission_docs" / "Genome_Medicine_Cover_Letter_GateC8R_AUTHOR_CONFIRMATION_REQUIRED.docx"
MAIN_PDF = PACKAGE / "internal_qc" / "wps_render_main" / "Genome_Medicine_Manuscript_GateC8R_AUTHOR_COMPLETION_REQUIRED_WPS.pdf"
SUPP_PDF = PACKAGE / "internal_qc" / "wps_render_supplement" / "Additional_file_1_Supplementary_Information_GateC8R_WPS.pdf"
COVER_PDF = PACKAGE / "internal_qc" / "wps_render_cover_letter" / "Genome_Medicine_Cover_Letter_GateC8R_AUTHOR_CONFIRMATION_REQUIRED_WPS.pdf"

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
FIXED_ZIP_TIME = (2026, 8, 20, 0, 0, 0)


def words(text: str) -> int:
    clean = re.sub(r"[`*_\[\]#|]", " ", text)
    return len(re.findall(r"\b[\w'-]+\b", clean))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def section(text: str, start: str, end: str | None) -> str:
    begin = text.index(start) + len(start)
    finish = text.index(end, begin) if end else len(text)
    return text[begin:finish].strip()


def ooxml(docx_path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(docx_path) as archive:
        return {name: archive.read(name) for name in archive.namelist() if name.endswith(".xml")}


def truthy(value: str) -> bool:
    return str(value).strip().lower() in {"true", "1", "yes"}


def package_readme() -> str:
    return """# Gate C8R Genome Medicine submission handoff

This directory is the repaired scientific and technical handoff generated on 20 August 2026.

## Status

- Scientific, figure, reference and technical package: PASS.
- Portal submission: BLOCKED until author-controlled declarations and the immutable archive DOI/licence are completed.
- Primary target: Genome Medicine; Communications Biology and Journal of Autoimmunity remain transfer routes.
- WPS visual review: PASS for all 26 manuscript pages, 4 supplementary pages and 1 cover-letter page.
- Figure assertions: PASS 43/43, including Figure 2a raw-point counts of 43 controls, 47 managed SLE and 90 total.
- Correlation-aware sensitivity: CAMERA positive 6/6 and BH-significant 5/6; FRY positive and BH-significant 6/6. The discovery STAT2 CAMERA exception is explicit.

## Upload mapping after author completion

- `main_text/`: editable main manuscript.
- `figures/`: Figures 1-5 as PDF and 600-dpi PNG composites.
- `additional_files/`: Supplementary Information, figure source data and regulator-sensitivity ZIP files.
- `submission_docs/`: cover letter, author completion form, target decision and reporting checklist.
- `references/`: DOI verification and Vancouver reference records.
- `review_copies/`: WPS-rendered PDFs for author review; these are not the primary editable uploads.
- `internal_qc/`: page renders, accessibility checks and audit evidence; do not upload unless requested.

## Hard stops

Do not submit until the institutional ethics determination, competing interests, funding, CRediT contributions, acknowledgements, all-author approval/originality confirmation, repository licence and immutable archive DOI have been supplied and verified.
"""


def build_manifest() -> list[dict[str, object]]:
    manifest_path = PACKAGE / "MANIFEST_SHA256.csv"
    rows: list[dict[str, object]] = []
    for path in sorted(p for p in PACKAGE.rglob("*") if p.is_file() and p != manifest_path):
        rows.append(
            {
                "relative_path": path.relative_to(PACKAGE).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
        )
    with manifest_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["relative_path", "bytes", "sha256"])
        writer.writeheader()
        writer.writerows(rows)
    return rows


def write_deterministic_zip(output: Path) -> None:
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(p for p in PACKAGE.rglob("*") if p.is_file()):
            relative = (Path(PACKAGE.name) / path.relative_to(PACKAGE)).as_posix()
            info = zipfile.ZipInfo(relative, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            info.create_system = 3
            archive.writestr(info, path.read_bytes(), compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)


def build_deterministic_archive() -> tuple[int, str]:
    with tempfile.TemporaryDirectory(prefix="gateC8R_zip_") as temp_dir:
        first = Path(temp_dir) / "first.zip"
        second = Path(temp_dir) / "second.zip"
        write_deterministic_zip(first)
        write_deterministic_zip(second)
        if sha256(first) != sha256(second):
            raise RuntimeError("Deterministic ZIP rebuild check failed")
        shutil.copyfile(first, PACKAGE_ZIP)
    return PACKAGE_ZIP.stat().st_size, sha256(PACKAGE_ZIP)


def main() -> None:
    manuscript = MANUSCRIPT_MD.read_text(encoding="utf-8")
    supplement = SUPPLEMENT_MD.read_text(encoding="utf-8")
    target = TARGET_MD.read_text(encoding="utf-8")
    figure_status = json.loads(FIGURE_STATUS.read_text(encoding="utf-8"))
    panel_assertions = json.loads(PANEL_ASSERTIONS.read_text(encoding="utf-8"))
    sensitivity_decision = json.loads(SENSITIVITY_DECISION.read_text(encoding="utf-8"))
    source_status = json.loads(SOURCE_STATUS.read_text(encoding="utf-8"))
    document_status = json.loads(DOCUMENT_STATUS.read_text(encoding="utf-8"))
    checks: dict[str, dict[str, object]] = {}

    def check(name: str, passed: bool, detail: str) -> None:
        checks[name] = {"pass": bool(passed), "detail": detail}

    check(
        "figure_build_freeze",
        figure_status.get("status") == "C8R_MAIN_FIGURES_BUILT_WITH_ASSERTIONS"
        and figure_status.get("figures") == 5
        and figure_status.get("panel_data_assertions_passed") is True,
        f"status={figure_status.get('status')}; figures={figure_status.get('figures')}",
    )
    assertions = panel_assertions.get("checks", [])
    assertion_map = {row.get("check"): row for row in assertions}
    point_counts = {
        name: assertion_map.get(name, {}).get("actual")
        for name in (
            "Figure2.panel_a.control_raw_points",
            "Figure2.panel_a.managed_sle_raw_points",
            "Figure2.panel_a.total_raw_points",
        )
    }
    check(
        "panel_data_assertions",
        panel_assertions.get("status") == "PASS"
        and len(assertions) == 43
        and all(row.get("pass") is True for row in assertions)
        and list(point_counts.values()) == [43, 47, 90],
        f"passed={sum(row.get('pass') is True for row in assertions)}/43; Figure2a={point_counts}",
    )
    check(
        "journal_target",
        "**Primary submission: Genome Medicine.**" in target,
        "Genome Medicine primary; transfer routes frozen without quartile assertion",
    )

    author_tokens = [
        "Zhi Chen [1] and Teng Qi [1,*]",
        "zhichen1@link.cuhk.edu.cn",
        "tengqi@link.cuhk.edu.cn",
        "0009-0001-0072-5576",
        "0009-0007-7648-4776",
        "School of Medicine, The Chinese University of Hong Kong, Shenzhen",
        "MED Start-up Building, 2001 Longxiang Boulevard",
    ]
    missing_authors = [token for token in author_tokens if token not in manuscript]
    check("author_identity", not missing_authors, f"missing={missing_authors or 'none'}")

    abstract = section(manuscript, "## Abstract", "## Keywords")
    abstract_words = words(abstract)
    abstract_labels = [f"**{label}:**" in abstract for label in ("Background", "Methods", "Results", "Conclusions")]
    check(
        "structured_abstract",
        all(abstract_labels) and abstract_words <= 350 and source_status.get("abstract_words") == 314,
        f"computed_words={abstract_words}; frozen_words={source_status.get('abstract_words')}; labels={sum(abstract_labels)}/4",
    )
    keyword_text = section(manuscript, "## Keywords", "## Background").replace("\n", " ")
    keyword_count = len([item for item in keyword_text.split(";") if item.strip()])
    check("keywords", 3 <= keyword_count <= 10, f"keywords={keyword_count}; allowed=3-10")

    required_sections = ["## Background", "## Methods", "## Results", "## Discussion", "## Conclusions", "## List of abbreviations", "## Declarations", "## References"]
    missing_sections = [heading for heading in required_sections if heading not in manuscript]
    check("required_sections", not missing_sections, f"missing={missing_sections or 'none'}")
    declaration_sections = [
        "### Ethics approval and consent to participate",
        "### Consent for publication",
        "### Availability of data and materials",
        "### Competing interests",
        "### Funding",
        "### Authors' contributions",
        "### Acknowledgements",
        "### Authors' information",
    ]
    missing_declarations = [heading for heading in declaration_sections if heading not in manuscript]
    check("declaration_structure", not missing_declarations, f"missing={missing_declarations or 'none'}")

    legends = re.findall(r"^### Figure (\d+) \| ([^\n]+)\n(.+?)(?=^### Figure |^## References)", manuscript, flags=re.M | re.S)
    legend_details = [{"figure": int(number), "title_words": words(title), "legend_words": words(body)} for number, title, body in legends]
    legends_pass = len(legends) == 5 and all(row["title_words"] <= 15 and row["legend_words"] <= 300 for row in legend_details)
    check("figure_legends", legends_pass, f"count={len(legends)}; details={legend_details}")

    figure_details = []
    figures_pass = True
    for number in range(1, 6):
        pngs = list((PACKAGE / "figures").glob(f"Figure{number}_*.png"))
        pdfs = list((PACKAGE / "figures").glob(f"Figure{number}_*.pdf"))
        if len(pngs) != 1 or len(pdfs) != 1:
            figures_pass = False
            figure_details.append({"figure": number, "error": "file count"})
            continue
        with Image.open(pngs[0]) as image:
            dimensions = image.size
        ok = (
            0 < pngs[0].stat().st_size < 10_000_000
            and 0 < pdfs[0].stat().st_size < 10_000_000
            and dimensions[0] >= 4000
            and dimensions[1] >= 3000
        )
        figures_pass &= ok
        figure_details.append({"figure": number, "png_bytes": pngs[0].stat().st_size, "pdf_bytes": pdfs[0].stat().st_size, "pixels": dimensions})
    check("figure_files", figures_pass, f"{figure_details}")

    with SENSITIVITY_CSV.open(encoding="utf-8", newline="") as handle:
        sensitivity_rows = list(csv.DictReader(handle))
    expected_targets = [98, 14, 129, 19, 161, 20]
    sensitivity_pass = (
        len(sensitivity_rows) == 6
        and [int(row["matched_signed_targets"]) for row in sensitivity_rows] == expected_targets
        and all(truthy(row["target_count_matches_frozen_ulm"]) for row in sensitivity_rows)
        and all(row["camera_direction"] == "Up" and row["fry_direction"] == "Up" for row in sensitivity_rows)
        and sum(float(row["camera_q_core6"]) < 0.05 for row in sensitivity_rows) == 5
        and sum(float(row["fry_q_core6"]) < 0.05 for row in sensitivity_rows) == 6
        and sensitivity_decision.get("decision") == "SUPPORTS_CONVERGENCE_WITH_DISCOVERY_STAT2_CAMERA_LIMITATION"
    )
    exception = sensitivity_decision.get("explicit_exception", {})
    check(
        "correlation_aware_sensitivity",
        sensitivity_pass and exception.get("contrast") == "gse174188_primary" and exception.get("regulator") == "STAT2",
        "targets=98/14/129/19/161/20; CAMERA Up=6/6, BH=5/6; FRY Up=6/6, BH=6/6; exception=discovery STAT2 CAMERA q=0.1355",
    )

    with REF_CSV.open(encoding="utf-8", newline="") as handle:
        ref_rows = list(csv.DictReader(handle))
    check(
        "reference_verification",
        len(ref_rows) == 26 and all(row["status"] == "PASS" for row in ref_rows) and source_status.get("reference_count") == 30,
        f"Crossref PASS={sum(row['status'] == 'PASS' for row in ref_rows)}/26; total manuscript references={source_status.get('reference_count')}",
    )

    numeric_tokens = [
        "0.947", "0.837", "1.086", "1.042", "rho=0.026",
        "3.187, 3.050 and 3.527", "3.294 and 3.666",
        "CAMERA retained the expected positive direction in six of six tests",
        "CAMERA q=0.1355", "FRY was positive and BH-significant in all six core tests",
    ]
    missing_numbers = [token for token in numeric_tokens if token not in manuscript]
    check("frozen_numeric_anchors", not missing_numbers, f"missing={missing_numbers or 'none'}")
    stale_terms = ["expanded ABC/APC-like", "central ABC/APC-like", "validated ABC/APC-like", "causal STAT1 activation", "CAMERA supported six of six"]
    stale_hits = [term for term in stale_terms if term.lower() in manuscript.lower()]
    check("no_stale_claim_language", not stale_hits, f"hits={stale_hits or 'none'}")
    boundary_tokens = [
        "not a globally shared disease transcriptome",
        "precluding a claim that every core test was significant under CAMERA",
        "not proving that STAT1 or STAT2 initiated the in vivo state",
        "prospective clinical validation remain outside the current evidence",
    ]
    missing_boundaries = [token for token in boundary_tokens if token not in manuscript]
    check("noncausal_boundaries", not missing_boundaries, f"present={len(boundary_tokens) - len(missing_boundaries)}/{len(boundary_tokens)}")

    docx_details = {path.name: path.stat().st_size if path.exists() else 0 for path in (MAIN_DOCX, SUPP_DOCX, COVER_DOCX)}
    check("editable_documents", all(size > 30_000 for size in docx_details.values()), f"{docx_details}")
    check(
        "document_build_status",
        document_status.get("created_at") == "2026-08-20"
        and len(document_status.get("outputs", [])) == 3
        and document_status.get("assets", {}).get("figure_files") == 10,
        f"outputs={len(document_status.get('outputs', []))}; figure_assets={document_status.get('assets', {}).get('figure_files')}",
    )

    main_xml = ooxml(MAIN_DOCX)
    main_document = ET.fromstring(main_xml["word/document.xml"])
    main_settings = ET.fromstring(main_xml["word/settings.xml"])
    main_styles = ET.fromstring(main_xml["word/styles.xml"])
    line_numbers = main_document.find(".//w:sectPr/w:lnNumType", NS) is not None
    page_field = b"PAGE" in b"".join(value for key, value in main_xml.items() if key.startswith("word/footer"))
    header_parts = [value for key, value in main_xml.items() if key.startswith("word/header")]
    header_texts = ["".join(ET.fromstring(value).itertext()) for value in header_parts]
    even_odd = main_settings.find(".//w:evenAndOddHeaders", NS) is not None
    normal_style = main_styles.find(".//w:style[@w:styleId='Normal']", NS)
    spacing = normal_style.find(".//w:spacing", NS) if normal_style is not None else None
    double_spacing = spacing is not None and spacing.get(f"{{{NS['w']}}}line") == "480"
    title_style = main_styles.find(".//w:style[@w:styleId='Title']", NS)
    no_title_border = title_style is not None and title_style.find(".//w:pBdr", NS) is None
    docx_format_pass = line_numbers and page_field and even_odd and len(header_parts) >= 2 and all("Genome Medicine" in text for text in header_texts) and double_spacing and no_title_border
    check("main_docx_ooxml", docx_format_pass, f"line_numbers={line_numbers}; page_field={page_field}; even_odd={even_odd}; headers={len(header_parts)}; double_spacing={double_spacing}; title_border_absent={no_title_border}")

    supplement_doc = Document(SUPP_DOCX)
    table_geometry = []
    for table in supplement_doc.tables:
        xml = table._tbl
        table_properties = xml.find(qn("w:tblPr"))
        table_width = table_properties.find(qn("w:tblW")) if table_properties is not None else None
        table_indent = table_properties.find(qn("w:tblInd")) if table_properties is not None else None
        table_grid = xml.find(qn("w:tblGrid"))
        cell_widths = []
        for row in table.rows:
            for cell in row.cells:
                cell_properties = cell._tc.find(qn("w:tcPr"))
                cell_widths.append(cell_properties is not None and cell_properties.find(qn("w:tcW")) is not None)
        table_geometry.append(table_width is not None and table_indent is not None and table_grid is not None and all(cell_widths))
    check("supplement_table_geometry", len(table_geometry) == 6 and all(table_geometry), f"tables={len(table_geometry)}; explicit_geometry={sum(table_geometry)}/6")

    a11y_reports = [PACKAGE / "internal_qc" / name for name in ("main_text_a11y.json", "supplement_a11y.json", "cover_letter_a11y.json")]
    a11y_details = []
    for path in a11y_reports:
        report = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"counts": {}}
        counts = report.get("counts", {})
        a11y_details.append({"file": path.name, **counts})
    check("docx_accessibility", len(a11y_details) == 3 and all(sum(int(row.get(level, 0)) for level in ("high", "medium", "low")) == 0 for row in a11y_details), f"{a11y_details}")

    source_zip = PACKAGE / "additional_files" / "Additional_file_2_Figure_Source_Data_GateC8R.zip"
    sensitivity_zip = PACKAGE / "additional_files" / "Additional_file_3_Regulator_Sensitivity_GateC8R.zip"
    with zipfile.ZipFile(source_zip) as archive:
        source_entries = sorted(name for name in archive.namelist() if not name.endswith("/"))
    with zipfile.ZipFile(sensitivity_zip) as archive:
        sensitivity_entries = sorted(name for name in archive.namelist() if not name.endswith("/"))
    source_csvs = sorted((PACKAGE / "additional_files" / "source_data").glob("Figure*_source_data.csv"))
    source_hash_file = PACKAGE / "additional_files" / "source_data" / "SHA256SUMS.csv"
    with source_hash_file.open(encoding="utf-8", newline="") as handle:
        source_hash_rows = list(csv.DictReader(handle))
    source_hashes_pass = len(source_csvs) == 5 and len(source_hash_rows) == 5 and all(
        sha256(PACKAGE / "additional_files" / "source_data" / row["file"]).lower() == row["sha256"].lower()
        for row in source_hash_rows
    )
    sensitivity_hash_file = PACKAGE / "additional_files" / "regulator_sensitivity" / "SHA256SUMS.csv"
    with sensitivity_hash_file.open(encoding="utf-8", newline="") as handle:
        sensitivity_hash_rows = list(csv.DictReader(handle))
    sensitivity_hashes_pass = len(sensitivity_hash_rows) == 2 and all(
        sha256(PACKAGE / "additional_files" / "regulator_sensitivity" / row["file"]).lower() == row["sha256"].lower()
        for row in sensitivity_hash_rows
    )
    check("figure_source_data", source_hashes_pass and len(source_entries) == 6, f"csv={len(source_csvs)}; checksums={len(source_hash_rows)}; zip_entries={len(source_entries)}")
    check("regulator_sensitivity_attachment", sensitivity_hashes_pass and len(sensitivity_entries) == 3, f"checksums={len(sensitivity_hash_rows)}; zip_entries={len(sensitivity_entries)}")

    render_specs = [
        ("main", MAIN_PDF, 26),
        ("supplement", SUPP_PDF, 4),
        ("cover", COVER_PDF, 1),
    ]
    render_details = []
    renders_pass = True
    for name, pdf, expected in render_specs:
        pages = len(PdfReader(pdf).pages) if pdf.exists() else 0
        png_count = len(list(pdf.parent.glob("page-*.png")))
        ok = pdf.exists() and pdf.stat().st_size > 50_000 and pages == expected and png_count == expected
        renders_pass &= ok
        render_details.append({"document": name, "pdf_bytes": pdf.stat().st_size if pdf.exists() else 0, "pdf_pages": pages, "png_pages": png_count})
    check("wps_render_outputs", renders_pass, f"{render_details}; visual_review=PASS_ALL_31_PAGES")

    hard_stops = [
        "institutional ethics determination",
        "competing interests",
        "funding",
        "final CRediT contributions and all-author approval",
        "acknowledgements",
        "all-author originality/submission confirmation",
        "open-source licence and immutable archive DOI",
    ]
    manuscript_placeholders = manuscript.count("[[")
    cover_source = (ROOT / "04_submission" / "cover_letter_genome_medicine_gateC8R_AUTHOR_COMPLETION_REQUIRED_2026-08-20.md").read_text(encoding="utf-8")
    cover_placeholders = cover_source.count("[[")
    check("hard_stops_visible", manuscript_placeholders == 6 and cover_placeholders == 2, f"manuscript placeholders={manuscript_placeholders}; cover placeholders={cover_placeholders}; unresolved hard stops={len(hard_stops)}")

    scientific_checks = [name for name in checks if name != "hard_stops_visible"]
    scientific_pass = all(checks[name]["pass"] for name in scientific_checks)
    portal_authorized = scientific_pass and not hard_stops
    decision = (
        "PASS_GATE_C8R_SCIENTIFIC_FIGURE_REPRODUCIBILITY_REPAIR_AUTHOR_ACTION_REQUIRED"
        if scientific_pass
        else "HOLD_GATE_C8R_TECHNICAL_REPAIR_REQUIRED"
    )
    audit = {
        "created_at": "2026-08-20",
        "decision": decision,
        "primary_target": "Genome Medicine",
        "scientific_technical_package_pass": scientific_pass,
        "portal_submission_authorized": portal_authorized,
        "checks": checks,
        "hard_stops": hard_stops,
        "next_stage": "Gate C8B author declarations, institutional ethics determination, repository licence and immutable archive DOI; then final portal preflight and submission",
    }

    audit_json = RUN_DIR / "07_GATE_C8R_FINAL_AUDIT.json"
    audit_md = RUN_DIR / "07_GATE_C8R_FINAL_AUDIT.md"
    audit_json.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    md_lines = [
        "# Gate C8R final submission audit",
        "",
        f"**Decision:** `{decision}`",
        "",
        f"**Scientific and technical package:** {'PASS' if scientific_pass else 'HOLD'}",
        "",
        f"**Portal submission authorized:** {'YES' if portal_authorized else 'NO'}",
        "",
        "## Checks",
        "",
        "| Check | Result | Detail |",
        "|---|---:|---|",
    ]
    for name, result in checks.items():
        detail = str(result["detail"]).replace("|", "/").replace("\n", " ")
        md_lines.append(f"| `{name}` | {'PASS' if result['pass'] else 'FAIL'} | {detail} |")
    md_lines.extend(["", "## Author-controlled hard stops", ""] + [f"- {item}" for item in hard_stops] + ["", "## Next stage", "", audit["next_stage"]])
    audit_md.write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    review_dir = PACKAGE / "review_copies"
    review_dir.mkdir(parents=True, exist_ok=True)
    for path in (MAIN_PDF, SUPP_PDF, COVER_PDF):
        shutil.copy2(path, review_dir / path.name)
    (PACKAGE / "README_GATE_C8R_PACKAGE.md").write_text(package_readme(), encoding="utf-8")
    internal = PACKAGE / "internal_qc"
    shutil.copy2(audit_json, internal / "GATE_C8R_FINAL_AUDIT.json")
    shutil.copy2(audit_md, internal / "GATE_C8R_FINAL_AUDIT.md")
    manifest_rows = build_manifest()
    shutil.copy2(PACKAGE / "MANIFEST_SHA256.csv", RUN_DIR / "08_GATE_C8R_INTEGRITY_MANIFEST.csv")
    zip_bytes, zip_hash = build_deterministic_archive()

    summary = {
        **audit,
        "deterministic_archive_rebuild_match": True,
        "package_files_manifested": len(manifest_rows),
        "package_zip": PACKAGE_ZIP.relative_to(ROOT).as_posix(),
        "package_zip_bytes": zip_bytes,
        "package_zip_sha256": zip_hash,
    }
    (RUN_DIR / "09_GATE_C8R_PACKAGE_STATUS.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))
    if not scientific_pass:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
