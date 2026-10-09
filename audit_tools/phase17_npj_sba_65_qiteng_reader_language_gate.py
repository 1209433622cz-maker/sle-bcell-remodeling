#!/usr/bin/env python3
"""Rebuild the narrow reader-language micro-gate from canonical sources."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "phase17_v7/npj_sba_qiteng_round3/20260929_reader_language_hard_negative"
INPUT_PACKAGE = RUN / "input_r2/SLE_Bcell_Submission_Package"
PACKAGE = RUN / "candidate/SLE_Bcell_Submission_Package"
ARTWORK = RUN / "artwork"
SUPPLEMENT_SOURCE = ROOT / "01_manuscript/Supplementary_Information.md"
CANONICAL_FIGURES = (
    ROOT
    / "phase17_v7/npj_sba_scientific_presentation_maintenance_freeze"
    / "20260908_final_cross_document/figures/figures"
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def apply_s7_layout(path: Path) -> dict[str, object]:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    document = Document(path)
    table = next(table for table in document.tables if table.cell(0, 0).text == "Result family")
    widths = [165.1 * ratio / 130 for ratio in (19, 27, 24, 19, 23, 18)]
    table.autofit = False
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Mm(width)
        for cell in column.cells:
            cell.width = Mm(width)
            properties = cell._tc.get_or_add_tcPr()
            margins = properties.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                properties.append(margins)
            for side in ("left", "right", "start", "end"):
                element = OxmlElement("w:" + side)
                element.set(qn("w:w"), "50")
                element.set(qn("w:type"), "dxa")
                margins.append(element)
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Supplementary Figure S2 |"):
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.space_before = Pt(0)
    document.core_properties.author = "Zhi Chen; Teng Qi"
    document.core_properties.subject = "QiTeng reader-language micro-gate"
    document.core_properties.comments = (
        "Reader-facing terminology synchronized without changing scientific data, tests or figures."
    )
    document.save(path)
    return {"s7_column_widths_mm": widths, "horizontal_padding_twips": 50}


def build(*, figures: bool, supplement: bool, initialize: bool) -> None:
    if not INPUT_PACKAGE.exists():
        raise FileNotFoundError(INPUT_PACKAGE)
    PACKAGE.parent.mkdir(parents=True, exist_ok=True)
    if initialize or not PACKAGE.exists():
        shutil.copytree(INPUT_PACKAGE, PACKAGE, dirs_exist_ok=True)

    status_path = RUN / "build_status.json"
    status = json.loads(status_path.read_text(encoding="utf-8")) if status_path.exists() else {}

    if figures:
        import phase17_c7_01_build_main_figures as main_figures

        figure_dir = ARTWORK / "figures"
        source_dir = ARTWORK / "source_data"
        figure_dir.mkdir(parents=True, exist_ok=True)
        source_dir.mkdir(parents=True, exist_ok=True)
        main_figures.ASSERTIONS.clear()
        main_figures.configure_style()
        main_figures.build_figure2(ROOT, figure_dir, source_dir)
        main_figures.build_figure3(ROOT, figure_dir, source_dir)

        data_archive = PACKAGE / "Supplementary_Data/Supplementary_Data_1_Figure_Source_Data.zip"
        source_checks: dict[str, bool] = {}
        with zipfile.ZipFile(data_archive) as archive:
            for number in (2, 3):
                name = f"Figure{number}_source_data.csv"
                source_checks[name] = (source_dir / name).read_bytes() == archive.read(name)
        if not all(source_checks.values()):
            raise RuntimeError(f"Figure source data changed: {source_checks}")

        figure_map = {
            "Figure2_sample_level_composition.pdf": PACKAGE / "Figures/Figure_2.pdf",
            "Figure3_gse174188_bconv_transcription.pdf": PACKAGE / "Figures/Figure_3.pdf",
        }
        parent_figure_hashes = {
            target.name: sha256(INPUT_PACKAGE / "Figures" / target.name)
            for target in figure_map.values()
        }
        for source_name, target in figure_map.items():
            shutil.copyfile(figure_dir / source_name, target)

        supplementary_output = ARTWORK / "supplementary_overlap_depletion"
        environment = dict(__import__("os").environ)
        environment["NPJ_SBA_STYLE"] = "1"
        subprocess.run(
            [
                sys.executable,
                str(ROOT / "audit_tools/phase17_round6_02_build_overlap_depletion_figure.py"),
                "--output-dir",
                str(supplementary_output),
            ],
            check=True,
            env=environment,
        )
        supplementary_figure_dir = ARTWORK / "supplementary_figures"
        shutil.copytree(CANONICAL_FIGURES, supplementary_figure_dir, dirs_exist_ok=True)
        for suffix in ("pdf", "png"):
            shutil.copyfile(
                supplementary_output / "figures" / f"Supplementary_Figure_S8_overlap_depletion.{suffix}",
                supplementary_figure_dir / f"Supplementary_Figure_S10_overlap_depletion.{suffix}",
            )
        with zipfile.ZipFile(data_archive) as archive:
            s10_source_matches = (
                supplementary_output / "source_data/Supplementary_Figure_S8_source_data.csv"
            ).read_bytes() == archive.read("Supplementary_Figure_S10_source_data.csv")
        if not s10_source_matches:
            raise RuntimeError("Supplementary Figure S10 source data changed")
        status.update({
            "source_data_byte_identical": source_checks,
            "figure_parent_sha256": parent_figure_hashes,
            "figure_candidate_sha256": {
                target.name: sha256(target) for target in figure_map.values()
            },
            "figure_assertions": len(main_figures.ASSERTIONS),
            "figure_assertions_passed": all(item["pass"] for item in main_figures.ASSERTIONS),
            "supplementary_figure_s10_source_data_byte_identical": s10_source_matches,
            "supplementary_figure_s10_sha256": sha256(
                supplementary_figure_dir / "Supplementary_Figure_S10_overlap_depletion.pdf"
            ),
        })
        (RUN / "figure_assertions.json").write_text(
            json.dumps(main_figures.ASSERTIONS, indent=2) + "\n", encoding="utf-8"
        )

    if supplement:
        import phase17_npj_sba_60_build_final_cross_document_documents as document_builder

        supplement_path = PACKAGE / "Supplementary_Information/Supplementary_Information.docx"
        document_builder.documents.markdown_to_docx(
            SUPPLEMENT_SOURCE,
            supplement_path,
            body_size=10.5,
            double_space=False,
            line_numbers=False,
            running_header="Supplementary information",
            title_override="Supplementary information",
            supplementary_figure_dirs=[ARTWORK / "supplementary_figures"],
            page_break_before_headings=set(),
        )
        document_builder.patch_supplement(supplement_path)
        layout = apply_s7_layout(supplement_path)
        status.update({
            "supplement_docx_sha256": sha256(supplement_path),
            "supplement_source": str(SUPPLEMENT_SOURCE.relative_to(ROOT)),
            "s7_layout": layout,
        })

    status.update({
        "status": "READER_LANGUAGE_SOURCE_REBUILD_COMPLETE_RENDER_REQUIRED",
        "manuscript_changed_in_this_gate": False,
        "statistical_outputs_changed": False,
    })
    status_path.write_text(json.dumps(status, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(status, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--figures", action="store_true")
    parser.add_argument("--supplement", action="store_true")
    parser.add_argument("--initialize", action="store_true")
    arguments = parser.parse_args()
    if not arguments.figures and not arguments.supplement:
        parser.error("select --figures and/or --supplement")
    build(figures=arguments.figures, supplement=arguments.supplement, initialize=arguments.initialize)
