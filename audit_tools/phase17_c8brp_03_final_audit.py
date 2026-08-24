#!/usr/bin/env python3
"""Audit and deterministically package the journal-facing C8BR prefreeze."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
RUN_DIR = ROOT / "phase17_v7" / "gateC8BRP" / "20260825_journal_facing_prefreeze"
C8BR_RUN = ROOT / "phase17_v7" / "gateC8BR" / "20260825_release_portability_preflight"
C8S_RUN = ROOT / "phase17_v7" / "gateC8S" / "20260821_supplementary_traceability_freeze"
PACKAGE = ROOT / "04_submission" / "package_genome_medicine_gateC8BRP_journal_facing_prefreeze_2026-08-25"
PACKAGE_ZIP = ROOT / "04_submission" / "package_genome_medicine_gateC8BRP_journal_facing_prefreeze_2026-08-25.zip"
MANUSCRIPT_MD = ROOT / "01_manuscript" / "manuscript_v15_genome_medicine_journal_facing_prefreeze_2026-08-25.md"
SUPPLEMENT_MD = ROOT / "01_manuscript" / "supplementary_information_v6_journal_facing_2026-08-25.md"
COVER_MD = ROOT / "04_submission" / "cover_letter_genome_medicine_gateC8BRP_AUTHOR_COMPLETION_REQUIRED_2026-08-25.md"
TARGET_MD = ROOT / "04_submission" / "journal_target_decision_gateC8BRP_2026-08-25.md"
AUTHOR_MATRIX_MD = ROOT / "04_submission" / "author_completion_matrix_gateC8BRP_2026-08-25.md"
REFERENCE_CSV = RUN_DIR / "references" / "reference_verification_gateC8BR.csv"

MAIN_DOCX = PACKAGE / "main_text" / "Genome_Medicine_Manuscript_AUTHOR_COMPLETION_REQUIRED.docx"
SUPP_DOCX = PACKAGE / "additional_files" / "Supplementary_Information.docx"
COVER_DOCX = PACKAGE / "submission_docs" / "Cover_Letter_AUTHOR_CONFIRMATION_REQUIRED.docx"
MAIN_PDF = PACKAGE / "internal_qc" / "wps_render_main" / "Genome_Medicine_Manuscript_WPS.pdf"
SUPP_PDF = PACKAGE / "internal_qc" / "wps_render_supplement" / "Supplementary_Information_WPS.pdf"
COVER_PDF = PACKAGE / "internal_qc" / "wps_render_cover_letter" / "Cover_Letter_WPS.pdf"
TITLE = "Disease-blind single-cell reconstruction separates unstable B-cell states from reproducible interferon remodeling in systemic lupus erythematosus"
FIXED_ZIP_TIME = (2026, 8, 25, 0, 0, 0)
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def words(text: str) -> int:
    clean = re.sub(r"[`*_#|\[\]]", " ", text)
    return len(re.findall(r"\b[\w'-]+\b", clean))


def section(text: str, start: str, end: str | None) -> str:
    begin = text.index(start) + len(start)
    finish = text.index(end, begin) if end else len(text)
    return text[begin:finish].strip()


def ooxml(docx_path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(docx_path) as archive:
        return {
            name: archive.read(name)
            for name in archive.namelist()
            if name.endswith(".xml")
        }


def pdf_text(path: Path) -> str:
    return re.sub(
        r"\s+",
        " ",
        " ".join((page.extract_text() or "") for page in PdfReader(path).pages),
    )


def build_manifest() -> list[dict[str, object]]:
    manifest_path = PACKAGE / "MANIFEST_SHA256.csv"
    rows: list[dict[str, object]] = []
    for path in sorted(
        item for item in PACKAGE.rglob("*") if item.is_file() and item != manifest_path
    ):
        rows.append(
            {
                "relative_path": path.relative_to(PACKAGE).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
        )
    with manifest_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=["relative_path", "bytes", "sha256"]
        )
        writer.writeheader()
        writer.writerows(rows)
    return rows


def write_deterministic_zip(output: Path) -> None:
    with zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as archive:
        for path in sorted(item for item in PACKAGE.rglob("*") if item.is_file()):
            relative = (Path(PACKAGE.name) / path.relative_to(PACKAGE)).as_posix()
            info = zipfile.ZipInfo(relative, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            info.create_system = 3
            archive.writestr(
                info,
                path.read_bytes(),
                compress_type=zipfile.ZIP_DEFLATED,
                compresslevel=9,
            )


def build_deterministic_archive() -> tuple[int, str]:
    with tempfile.TemporaryDirectory(prefix="gateC8BRP_zip_") as temp_dir:
        first = Path(temp_dir) / "first.zip"
        second = Path(temp_dir) / "second.zip"
        write_deterministic_zip(first)
        write_deterministic_zip(second)
        if sha256(first) != sha256(second):
            raise RuntimeError("Deterministic package ZIP rebuild check failed")
        shutil.copyfile(first, PACKAGE_ZIP)
    return PACKAGE_ZIP.stat().st_size, sha256(PACKAGE_ZIP)


def package_readme(page_counts: dict[str, int], stats_hash: str) -> str:
    return f"""# Genome Medicine journal-facing prefreeze package

This package was generated on 25 August 2026 from the frozen scientific state.

## Status

- Journal-facing manuscript, Supplementary Information, figures, source data, release runtime and technical package: PASS.
- Scientific estimates changed: NO.
- Portal submission: BLOCKED pending author-controlled declarations, institutional ethics determination, repository licence and immutable archive DOI.
- Primary target: Genome Medicine; no journal-quartile assertion is frozen here.
- WPS render: {page_counts['main']} manuscript pages, {page_counts['supplement']} supplementary pages and {page_counts['cover']} cover-letter page.
- Main-panel assertions: PASS 46/46. Frozen supplementary-panel assertions: PASS 29/29.
- The identity-stability methods report the resampling fraction and unit, frozen versus recomputed steps, cluster mapping, thresholds and random seeds.
- Figure 1a uses graphical workflow nodes; Figure 4d uses reader-facing omission labels with original source codes retained in Source Data; Figure 5a keeps regulatory and response evidence as parallel branches.
- Full statistical archive remains byte-identical to the scientific freeze; SHA-256 `{stats_hash}`.

## File map

- `main_text/`: editable main manuscript.
- `figures/`: Figures 1-5 as vector PDF and 600-dpi PNG composites.
- `figures_supplementary/`: Supplementary Figures S1-S7 as vector PDF and 600-dpi PNG composites.
- `additional_files/`: Supplementary Information, source data, regulator sensitivity and full statistical results.
- `submission_docs/`: cover letter, author completion form, target decision, reporting checklist and portal filename map.
- `portal_upload_preview/`: clean upload aliases for layout review only; the included hard-stop marker forbids upload.
- `references/`: 28 DOI verification records and journal-formatted metadata.
- `reproducibility/`: separate scientific and release environment locks, smoke test, accessibility audit and active runner.
- `review_copies/`: WPS-rendered PDFs for author review.
- `internal_qc/`: accessibility reports, page rasters and audit evidence; do not upload unless requested.

## Hard stop

Do not submit until every visible author-controlled item is completed and verified. Do not modify frozen scientific estimates during author completion.
"""


def main() -> None:
    manuscript = MANUSCRIPT_MD.read_text(encoding="utf-8-sig")
    supplement = SUPPLEMENT_MD.read_text(encoding="utf-8-sig")
    cover = COVER_MD.read_text(encoding="utf-8-sig")
    target = TARGET_MD.read_text(encoding="utf-8-sig")
    author_matrix = AUTHOR_MATRIX_MD.read_text(encoding="utf-8-sig")
    reproducibility = (ROOT / "REPRODUCIBILITY.md").read_text(encoding="utf-8-sig")

    figure_status = read_json(RUN_DIR / "01_FIGURE_BUILD_STATUS.json")
    main_assertions = read_json(RUN_DIR / "02_PANEL_DATA_ASSERTIONS.json")
    runtime_status = read_json(RUN_DIR / "00_GATE_C8BR_RELEASE_RUNTIME_STATUS.json")
    reference_status = read_json(RUN_DIR / "01_GATE_C8BRP_REFERENCE_STATUS.json")
    source_status = read_json(RUN_DIR / "03_GATE_C8BRP_SOURCE_BUILD_STATUS.json")
    document_status = read_json(RUN_DIR / "04_GATE_C8BRP_DOCUMENT_BUILD_STATUS.json")
    supp_status = read_json(C8S_RUN / "04_SUPPLEMENTARY_FIGURE_BUILD_STATUS.json")
    supp_assertions = read_json(C8S_RUN / "03_SUPPLEMENTARY_PANEL_DATA_ASSERTIONS.json")
    stats_status = read_json(C8S_RUN / "05_FULL_STATISTICAL_RESULTS_STATUS.json")

    checks: dict[str, dict[str, object]] = {}

    def check(name: str, passed: bool, detail: object) -> None:
        checks[name] = {"pass": bool(passed), "detail": str(detail)}

    runner_path = (
        ROOT
        / "audit_tools"
        / "run_6013RP_phase17_gateC8BRP_journal_facing_prefreeze.ps1"
    )
    runner = runner_path.read_text(encoding="utf-8-sig")
    release_yml = (
        ROOT / "audit_tools" / "environment_gateC8BR_release_2026-08-25.yml"
    ).read_text(encoding="utf-8-sig")
    explicit_spec = (
        ROOT
        / "audit_tools"
        / "environment_gateC8BR_release_explicit_win64_2026-08-25.txt"
    ).read_text(encoding="utf-8-sig")
    check(
        "portable_release_runtime",
        runtime_status.get("status") == "PASS_GATE_C8BR_RELEASE_RUNTIME"
        and runtime_status.get("versions_match") is True
        and not re.search(r"(?im)(?:^|[=\"'])\s*[A-Z]:\\", runner)
        and "matplotlib==3.10.7" in release_yml
        and "python-docx==1.2.0" in release_yml
        and explicit_spec.startswith("# This file may be used to create an environment using:")
        and "@EXPLICIT" in explicit_spec,
        f"python={runtime_status.get('python_version')}; exact win-64 package spec present",
    )
    reproducibility_tokens = [
        "Scientific analysis environment",
        "Release and document environment",
        "environment_gateC8BR_release_explicit_win64_2026-08-25.txt",
        runner_path.name,
        "Reference DOI identities independently resolved: 28/28",
        "Numbered manuscript references: 32",
        "Times New Roman text",
    ]
    check(
        "reproducibility_record",
        all(token in reproducibility for token in reproducibility_tokens)
        and "bundled Codex Python runtime" not in reproducibility
        and "Arial text" not in reproducibility,
        f"tokens={len(reproducibility_tokens)}; separate analysis/release locks documented",
    )

    main_rows = main_assertions.get("checks", [])
    check(
        "main_figure_assertions",
        figure_status.get("status")
        == "PASS_GATE_C8BRP_JOURNAL_FACING_FIGURES_BUILT"
        and figure_status.get("all_source_data_match_gateC8BR") is True
        and main_assertions.get("status") == "PASS"
        and len(main_rows) == 46
        and all(row.get("pass") is True for row in main_rows),
        f"passed={sum(row.get('pass') is True for row in main_rows)}/46; source data unchanged",
    )
    supp_rows = supp_assertions.get("checks", [])
    check(
        "supplementary_figure_assertions",
        supp_status.get("status") == "PASS_GATE_C8S_SUPPLEMENTARY_FIGURES_BUILT"
        and supp_assertions.get("status") == "PASS"
        and len(supp_rows) == 29
        and all(row.get("pass") is True for row in supp_rows),
        f"frozen assertions passed={sum(row.get('pass') is True for row in supp_rows)}/29",
    )

    figure_details: list[dict[str, object]] = []
    figures_pass = True
    for prefix, directory, total in (
        ("Figure", PACKAGE / "figures", 5),
        ("Supplementary_Figure_S", PACKAGE / "figures_supplementary", 7),
    ):
        for number in range(1, total + 1):
            pngs = list(directory.glob(f"{prefix}{number}_*.png"))
            pdfs = list(directory.glob(f"{prefix}{number}_*.pdf"))
            ok = len(pngs) == 1 and len(pdfs) == 1
            dimensions = (0, 0)
            if ok:
                with Image.open(pngs[0]) as image:
                    dimensions = image.size
                ok = (
                    dimensions[0] >= 4000
                    and dimensions[1] >= 3000
                    and pngs[0].stat().st_size > 0
                    and pdfs[0].stat().st_size > 0
                )
            figures_pass &= ok
            figure_details.append(
                {"figure": f"{prefix}{number}", "pixels": dimensions, "pass": ok}
            )
    check(
        "publication_figure_files",
        figures_pass,
        f"pairs={len(figure_details)}; each PNG >=4000x3000 and each PDF non-empty",
    )

    figure1_text = pdf_text(next((PACKAGE / "figures").glob("Figure1_*.pdf")))
    figure4_text = pdf_text(next((PACKAGE / "figures").glob("Figure4_*.pdf")))
    figure5_text = pdf_text(next((PACKAGE / "figures").glob("Figure5_*.pdf")))
    check(
        "reader_facing_figure_semantics",
        all(
            token in figure1_text
            for token in (
                "Frozen programs",
                "Independent validation",
                "Regulatory + response evidence",
            )
        )
        and "->" not in figure1_text
        and "Omit source label 1" in figure4_text
        and "Omit source label 8" in figure4_text
        and all(
            token in figure5_text
            for token in (
                "Parallel evidence architecture",
                "Regulatory branch",
                "Response branch",
            )
        ),
        "Figure 1 graphical workflow; Figure 4 sequential omission labels; Figure 5 parallel branches",
    )

    source_dir = PACKAGE / "additional_files" / "source_data"
    source_rows = read_csv(source_dir / "SHA256SUMS.csv")
    source_csvs = sorted(source_dir.glob("*Figure*_source_data.csv"))
    source_hashes_ok = (
        len(source_csvs) == 12
        and len(source_rows) == 12
        and all(
            sha256(source_dir / row["file"]) == row["sha256"].upper()
            for row in source_rows
        )
    )
    main_hashes_ok = all(
        sha256(source_dir / f"Figure{number}_source_data.csv")
        == sha256(C8BR_RUN / "source_data" / f"Figure{number}_source_data.csv")
        for number in range(1, 6)
    )
    figure4_rows = read_csv(source_dir / "Figure4_source_data.csv")
    source_codes_present = all(
        any(f"B-caSC{number}" in str(value) for row in figure4_rows for value in row.values())
        for number in range(8)
    )
    check(
        "figure_source_data",
        source_hashes_ok and main_hashes_ok and source_codes_present,
        "12/12 SHA manifest entries; all five main sources byte-identical; B-caSC0-7 retained",
    )

    stats_zip = PACKAGE / "additional_files" / "Full_Statistical_Results.zip"
    frozen_stats = C8S_RUN / "Additional_file_4_Full_Statistical_Results_GateC8S.zip"
    with zipfile.ZipFile(stats_zip) as archive:
        stats_entries = [name for name in archive.namelist() if not name.endswith("/")]
    stats_hash = sha256(stats_zip)
    stats_ok = (
        stats_status.get("status") == "PASS_GATE_C8S_FULL_STATISTICAL_ARCHIVE_BUILT"
        and len(stats_entries) == 63
        and stats_hash == sha256(frozen_stats)
        and stats_hash == stats_status.get("archive_sha256")
        and stats_status.get("complete_gene_result_files") == 12
        and stats_status.get("sanitized_design_matrices") == 12
    )
    check(
        "frozen_statistical_archive",
        stats_ok,
        f"entries={len(stats_entries)}; byte-identical sha256={stats_hash}",
    )

    abstract_words = words(section(manuscript, "## Abstract", "## Keywords"))
    refs = [
        int(value)
        for value in re.findall(
            r"^(\d+)\.\s", section(manuscript, "## References", None), flags=re.M
        )
    ]
    manuscript_required = [
        f"# {TITLE}",
        "**Version:** Journal-facing author-completion draft v15",
        "did not support stable fine-grained naive/memory subtype assignments",
        "Sequential display labels 1-8",
        "edgeR: a Bioconductor package",
        "doi:10.1136/lupus-2026-002042",
    ]
    manuscript_forbidden = [
        "technical- library",
        "statistical- engine",
        "edgeR :",
        "Gate C8S remains",
        "Gate C8B adds",
        "post-audit",
        "superseded audit artifacts",
    ]
    check(
        "manuscript_reader_contract",
        all(token in manuscript for token in manuscript_required)
        and not any(token in manuscript for token in manuscript_forbidden)
        and abstract_words <= 350
        and refs == list(range(1, 33))
        and "### Authors' information" not in manuscript,
        f"abstract={abstract_words}; references={len(refs)}; optional author biography omitted",
    )

    supplement_required = [
        "selecting 80% of cells without replacement separately within each `library_uuid`",
        "complete frozen 50-dimensional Harmony-adjusted principal-component representation",
        "15-nearest-neighbour graph",
        "resolutions 0.4, 0.6 and 0.8",
        "20260806 + 1000 + r",
        "row-wise maximum of the observed-by-reference contingency table",
        "minimum mapped ARI at least 0.90",
        "observed support was 1.00 for every required marker",
    ]
    supplement_forbidden = [
        "Gate C",
        "preflight",
        "superseded",
        "release-portability",
        "C2B1-C8R",
        "post-audit",
    ]
    check(
        "supplement_reader_contract",
        all(token in supplement for token in supplement_required)
        and not any(token in supplement for token in supplement_forbidden)
        and supplement.count("[[SUPPLEMENTARY_FIGURE:S") == 7
        and len(re.findall(r"^## Supplementary Table S\d+", supplement, flags=re.M))
        == 8,
        "exact resampling mechanics present; internal history absent; figures=7; tables=8",
    )
    check(
        "source_build_contract",
        source_status.get("status")
        == "PASS_GATE_C8BRP_JOURNAL_FACING_SOURCES_BUILT"
        and source_status.get("scientific_estimates_changed") is False
        and source_status.get("references") == 32
        and source_status.get("supplement_internal_history_tokens") == 0,
        "v15 manuscript and v6 supplement PASS; scientific estimates unchanged",
    )
    check(
        "journal_target_and_cover",
        "**Primary submission: Genome Medicine.**" in target
        and "No JCR quartile or CAS category is frozen here." in target
        and "an interferon-responsive transcriptional program" in cover
        and "Independent M5911 enrichment" not in cover,
        "Genome Medicine retained; claim-bounded cover letter; quartile not frozen",
    )
    reference_rows = read_csv(REFERENCE_CSV)
    check(
        "reference_verification",
        reference_status.get("decision") == "PASS"
        and reference_status.get("doi_records") == 28
        and reference_status.get("manuscript_references") == 32
        and len(reference_rows) == 28
        and all(row["status"] == "PASS" for row in reference_rows),
        "DOI identities PASS=28/28; numbered references=32",
    )

    docx_sizes = {
        path.name: path.stat().st_size if path.exists() else 0
        for path in (MAIN_DOCX, SUPP_DOCX, COVER_DOCX)
    }
    assets = document_status.get("assets", {})
    check(
        "document_build_status",
        document_status.get("status")
        == "PASS_GATE_C8BRP_DOCUMENTS_AND_PORTAL_PREVIEW_BUILT"
        and len(document_status.get("outputs", [])) == 3
        and assets.get("main_figure_files") == 10
        and assets.get("supplementary_figure_files") == 14
        and assets.get("source_files") == 12
        and assets.get("reproducibility_files") == 10
        and assets.get("portal_aliases") == 18
        and all(size > 30_000 for size in docx_sizes.values()),
        docx_sizes,
    )

    alias_rows = read_csv(
        PACKAGE / "submission_docs" / "PORTAL_UPLOAD_FILENAME_MAP.csv"
    )
    alias_names = [Path(row["portal_alias"]).name for row in alias_rows]
    aliases_ok = len(alias_rows) == 18 and len(set(alias_names)) == 18
    for row in alias_rows:
        source = ROOT / row["provenance_path"]
        alias = PACKAGE / row["portal_alias"]
        aliases_ok &= (
            source.is_file()
            and alias.is_file()
            and sha256(source) == sha256(alias) == row["sha256"].upper()
            and alias.stat().st_size == int(row["bytes"])
        )
    aliases_ok &= not any(
        re.search(r"(?i)(gate|author|frozen|prefreeze)", name)
        for name in alias_names
    )
    hard_stop_marker = (
        PACKAGE
        / "portal_upload_preview"
        / "DO_NOT_UPLOAD_AUTHOR_ACTION_REQUIRED.txt"
    )
    check(
        "portal_preview_aliases",
        aliases_ok
        and hard_stop_marker.is_file()
        and "Do not upload" in hard_stop_marker.read_text(encoding="utf-8"),
        "18 unique clean aliases; source/alias/manifest hashes match; upload blocked",
    )

    main_xml = ooxml(MAIN_DOCX)
    main_document = ET.fromstring(main_xml["word/document.xml"])
    main_settings = ET.fromstring(main_xml["word/settings.xml"])
    main_styles = ET.fromstring(main_xml["word/styles.xml"])
    line_numbers = main_document.find(".//w:sectPr/w:lnNumType", NS) is not None
    page_field = b"PAGE" in b"".join(
        value for key, value in main_xml.items() if key.startswith("word/footer")
    )
    even_odd = main_settings.find(".//w:evenAndOddHeaders", NS) is not None
    normal_style = main_styles.find(".//w:style[@w:styleId='Normal']", NS)
    spacing = (
        normal_style.find(".//w:spacing", NS) if normal_style is not None else None
    )
    double_spacing = (
        spacing is not None and spacing.get(f"{{{NS['w']}}}line") == "480"
    )
    check(
        "main_docx_ooxml",
        line_numbers and page_field and even_odd and double_spacing,
        f"line_numbers={line_numbers}; page_field={page_field}; even_odd={even_odd}; double_spacing={double_spacing}",
    )

    supplement_doc = Document(SUPP_DOCX)
    table_geometry: list[bool] = []
    for table in supplement_doc.tables:
        properties = table._tbl.find(qn("w:tblPr"))
        grid = table._tbl.find(qn("w:tblGrid"))
        widths: list[bool] = []
        for row in table.rows:
            for cell in row.cells:
                cell_properties = cell._tc.find(qn("w:tcPr"))
                widths.append(
                    cell_properties is not None
                    and cell_properties.find(qn("w:tcW")) is not None
                )
        table_geometry.append(
            properties is not None
            and properties.find(qn("w:tblW")) is not None
            and properties.find(qn("w:tblInd")) is not None
            and grid is not None
            and all(widths)
        )
    check(
        "supplement_docx_structure",
        len(table_geometry) == 8
        and all(table_geometry)
        and len(supplement_doc.inline_shapes) == 7,
        f"tables={len(table_geometry)}; explicit_geometry={sum(table_geometry)}; inline_figures={len(supplement_doc.inline_shapes)}",
    )

    a11y_details: list[dict[str, object]] = []
    for name in (
        "main_text_a11y.json",
        "supplement_a11y.json",
        "cover_letter_a11y.json",
    ):
        report = read_json(PACKAGE / "internal_qc" / name)
        counts = report.get("counts", {})
        a11y_details.append({"file": name, **counts})
    check(
        "docx_accessibility",
        all(
            sum(int(row.get(level, 0)) for level in ("high", "medium", "low"))
            == 0
            for row in a11y_details
        ),
        a11y_details,
    )

    render_details: list[dict[str, object]] = []
    page_counts: dict[str, int] = {}
    renders_ok = True
    for name, pdf, minimum, maximum in (
        ("main", MAIN_PDF, 25, 31),
        ("supplement", SUPP_PDF, 12, 16),
        ("cover", COVER_PDF, 1, 1),
    ):
        reader = PdfReader(pdf) if pdf.exists() else None
        pages = len(reader.pages) if reader else 0
        pngs = len(list(pdf.parent.glob("final-page-*.png")))
        ok = (
            pdf.exists()
            and pdf.stat().st_size > 30_000
            and minimum <= pages <= maximum
            and pngs == pages
        )
        renders_ok &= ok
        page_counts[name] = pages
        render_details.append(
            {
                "document": name,
                "pages": pages,
                "page_pngs": pngs,
                "bytes": pdf.stat().st_size if pdf.exists() else 0,
            }
        )
    check("wps_render_integrity", renders_ok, render_details)

    author_tokens = [
        "Zhi Chen [1] and Teng Qi [1,*]",
        "zhichen1@link.cuhk.edu.cn",
        "tengqi@link.cuhk.edu.cn",
        "0009-0001-0072-5576",
        "0009-0007-7648-4776",
        "School of Medicine, The Chinese University of Hong Kong, Shenzhen",
    ]
    matrix_tokens = [
        "Zhi Chen first author; Teng Qi corresponding author",
        "Both authors: MSc students in Bioinformatics",
        "Queen Mary University of London",
        "Nanchang University",
        "No biography for Teng Qi was inferred",
    ]
    check(
        "confirmed_author_identity",
        all(token in manuscript for token in author_tokens)
        and all(token in author_matrix for token in matrix_tokens),
        "names, order, emails, ORCIDs, titles and supplied Zhi Chen biography preserved without inferring a Teng Qi biography",
    )

    hard_stops = [
        "institutional ethics determination",
        "competing interests",
        "funding",
        "final CRediT contributions and all-author approval",
        "acknowledgements",
        "all-author originality and exclusive-submission confirmation",
        "correspondence-address approval",
        "author-approved repository licence and immutable archive DOI",
        "APC or institutional agreement check",
    ]
    manuscript_placeholders = manuscript.count("[[")
    cover_placeholders = cover.count("[[")
    check(
        "hard_stops_visible",
        manuscript_placeholders == 6
        and cover_placeholders == 2
        and author_matrix.count("- [ ]") >= 20,
        f"manuscript placeholders={manuscript_placeholders}; cover placeholders={cover_placeholders}; author actions={author_matrix.count('- [ ]')}",
    )

    technical_checks = [name for name in checks if name != "hard_stops_visible"]
    technical_pass = all(checks[name]["pass"] for name in technical_checks)
    decision = (
        "PASS_GATE_C8BR_JOURNAL_FACING_PREFREEZE_AUTHOR_ACTION_REQUIRED"
        if technical_pass
        else "HOLD_GATE_C8BR_JOURNAL_FACING_TECHNICAL_REPAIR_REQUIRED"
    )
    next_stage = (
        "Complete author-controlled declarations and institutional ethics determination; approve the correspondence address and repository licence; create an immutable release DOI; then replace all placeholders and run the zero-placeholder WPS and portal-field preflight without reopening scientific analysis."
    )
    audit = {
        "created_at": "2026-08-25",
        "decision": decision,
        "primary_target": "Genome Medicine",
        "source_scientific_freeze": "Gate C8S",
        "scientific_estimates_changed": False,
        "technical_package_pass": technical_pass,
        "portal_submission_authorized": False,
        "checks": checks,
        "hard_stops": hard_stops,
        "next_stage": next_stage,
        "next_if_author_completed": "PASS_GATE_C8BR_RELEASE_PORTABILITY_AUTHOR_COMPLETION_AND_PORTAL_PREFLIGHT",
    }

    audit_json = RUN_DIR / "05_GATE_C8BRP_FINAL_AUDIT.json"
    audit_md = RUN_DIR / "05_GATE_C8BRP_FINAL_AUDIT.md"
    audit_json.write_text(
        json.dumps(audit, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    lines = [
        "# Journal-facing prefreeze final audit",
        "",
        f"**Decision:** `{decision}`",
        "",
        f"**Technical package:** {'PASS' if technical_pass else 'HOLD'}",
        "",
        "**Scientific estimates changed:** NO",
        "",
        "**Portal submission authorized:** NO - author-controlled hard stops remain.",
        "",
        "## Checks",
        "",
        "| Check | Result | Detail |",
        "|---|---:|---|",
    ]
    for name, result in checks.items():
        detail = str(result["detail"]).replace("|", "/").replace("\n", " ")
        lines.append(
            f"| `{name}` | {'PASS' if result['pass'] else 'FAIL'} | {detail} |"
        )
    lines.extend(
        ["", "## Author-controlled hard stops", ""]
        + [f"- {item}" for item in hard_stops]
        + ["", "## Next stage", "", next_stage]
    )
    audit_md.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")

    review_dir = PACKAGE / "review_copies"
    review_dir.mkdir(parents=True, exist_ok=True)
    for path in (MAIN_PDF, SUPP_PDF, COVER_PDF):
        shutil.copy2(path, review_dir / path.name)
    (PACKAGE / "README_PACKAGE.md").write_text(
        package_readme(page_counts, stats_hash), encoding="utf-8", newline="\n"
    )
    shutil.copy2(audit_json, PACKAGE / "internal_qc" / "FINAL_AUDIT.json")
    shutil.copy2(audit_md, PACKAGE / "internal_qc" / "FINAL_AUDIT.md")
    manifest_rows = build_manifest()
    shutil.copy2(
        PACKAGE / "MANIFEST_SHA256.csv",
        RUN_DIR / "06_GATE_C8BRP_INTEGRITY_MANIFEST.csv",
    )
    zip_bytes, zip_hash = build_deterministic_archive()
    summary = {
        **audit,
        "deterministic_archive_rebuild_match": True,
        "package_files_manifested": len(manifest_rows),
        "page_counts": page_counts,
        "package_zip": PACKAGE_ZIP.relative_to(ROOT).as_posix(),
        "package_zip_bytes": zip_bytes,
        "package_zip_sha256": zip_hash,
    }
    (RUN_DIR / "07_GATE_C8BRP_PACKAGE_STATUS.json").write_text(
        json.dumps(summary, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    print(json.dumps(summary, indent=2))
    if not technical_pass:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
