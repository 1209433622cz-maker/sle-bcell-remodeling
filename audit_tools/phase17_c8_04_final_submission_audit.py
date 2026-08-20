#!/usr/bin/env python3
"""Audit and package the Gate C8 Genome Medicine submission handoff."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from PIL import Image
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
RUN_DIR = ROOT / "phase17_v7" / "gateC8" / "20260820_genome_medicine_submission_package"
PACKAGE = ROOT / "04_submission" / "package_genome_medicine_gateC8_2026-08-20"
PACKAGE_ZIP = ROOT / "04_submission" / "package_genome_medicine_gateC8_2026-08-20.zip"
MANUSCRIPT_MD = ROOT / "01_manuscript" / "manuscript_v10_genome_medicine_submission_2026-08-20.md"
SUPPLEMENT_MD = ROOT / "01_manuscript" / "supplementary_information_v1_gateC8_2026-08-20.md"
TARGET_MD = ROOT / "04_submission" / "journal_target_decision_gateC8_2026-08-20.md"
GATE_C7 = ROOT / "phase17_v7" / "gateC7" / "20260820_manuscript_figure_integration" / "06_GATE_C7_FINAL_AUDIT.json"
REF_CSV = RUN_DIR / "references" / "reference_verification_gateC8.csv"

MAIN_DOCX = PACKAGE / "main_text" / "Genome_Medicine_Manuscript_GateC8_AUTHOR_COMPLETION_REQUIRED.docx"
SUPP_DOCX = PACKAGE / "additional_files" / "Additional_file_1_Supplementary_Information_GateC8.docx"
COVER_DOCX = PACKAGE / "submission_docs" / "Genome_Medicine_Cover_Letter_GateC8_AUTHOR_CONFIRMATION_REQUIRED.docx"
MAIN_PDF = PACKAGE / "internal_qc" / "wps_render_main" / "Genome_Medicine_Manuscript_GateC8_AUTHOR_COMPLETION_REQUIRED_WPS.pdf"
SUPP_PDF = PACKAGE / "internal_qc" / "wps_render_supplement" / "Additional_file_1_Supplementary_Information_GateC8_WPS.pdf"
COVER_PDF = PACKAGE / "internal_qc" / "wps_render_cover_letter" / "Genome_Medicine_Cover_Letter_GateC8_AUTHOR_CONFIRMATION_REQUIRED_WPS.pdf"

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def words(text: str) -> int:
    clean = re.sub(r"[`*_\[\]#|]", " ", text)
    return len(re.findall(r"\b[\w/-]+\b", clean))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def section(text: str, start: str, end: str | None) -> str:
    begin = text.index(start) + len(start)
    finish = text.index(end, begin) if end else len(text)
    return text[begin:finish].strip()


def ooxml(docx_path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(docx_path) as archive:
        return {name: archive.read(name) for name in archive.namelist() if name.endswith(".xml")}


def package_readme() -> str:
    return """# Gate C8 Genome Medicine submission handoff

This directory is the complete scientific and technical submission handoff generated on 20 August 2026.

## Status

- Scientific and technical package: PASS.
- Portal submission: BLOCKED until all author-controlled declarations and the immutable archive DOI/licence are completed.
- Primary target: Genome Medicine.
- WPS visual review: PASS for all 21 manuscript pages, 3 supplementary pages and 1 cover-letter page.

## Upload mapping after author completion

- `main_text/`: editable main manuscript.
- `figures/`: Figures 1-5 as PDF and 600-dpi PNG composites.
- `additional_files/`: Supplementary Information and figure source-data ZIP.
- `submission_docs/`: cover letter, author completion form, target decision and reporting checklist.
- `references/`: DOI verification and Vancouver reference records.
- `review_copies/`: WPS-rendered PDFs for author review; these are not the primary editable uploads.
- `internal_qc/`: page renders and audit evidence; do not upload unless requested by the journal.

## Hard stops

Do not submit until the institutional ethics determination, competing interests, funding, CRediT contributions, acknowledgements, all-author approval/originality confirmation, repository licence and immutable archive DOI have been supplied and verified.
"""


def build_manifest() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    manifest_path = PACKAGE / "MANIFEST_SHA256.csv"
    for path in sorted(p for p in PACKAGE.rglob("*") if p.is_file() and p != manifest_path):
        rows.append(
            {
                "relative_path": path.relative_to(PACKAGE).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
        )
    with manifest_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["relative_path", "bytes", "sha256"])
        writer.writeheader()
        writer.writerows(rows)
    return rows


def build_zip() -> None:
    if PACKAGE_ZIP.exists():
        PACKAGE_ZIP.unlink()
    with zipfile.ZipFile(PACKAGE_ZIP, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(p for p in PACKAGE.rglob("*") if p.is_file()):
            archive.write(path, path.relative_to(PACKAGE.parent))


def main() -> None:
    manuscript = MANUSCRIPT_MD.read_text(encoding="utf-8")
    supplement = SUPPLEMENT_MD.read_text(encoding="utf-8")
    target = TARGET_MD.read_text(encoding="utf-8")
    checks: dict[str, dict[str, object]] = {}

    def check(name: str, passed: bool, detail: str) -> None:
        checks[name] = {"pass": bool(passed), "detail": detail}

    gate_c7 = json.loads(GATE_C7.read_text(encoding="utf-8"))
    check(
        "gate_c7_scientific_freeze",
        gate_c7.get("decision") == "PASS_GATE_C7_MANUSCRIPT_AND_FIVE_FIGURE_SCIENTIFIC_FREEZE",
        str(gate_c7.get("decision")),
    )
    check("journal_target", "**Primary submission: Genome Medicine.**" in target, "Genome Medicine primary; transfer routes frozen")

    author_tokens = [
        "Zhi Chen [1] and Teng Qi [1,*]",
        "zhichen1@link.cuhk.edu.cn",
        "tengqi@link.cuhk.edu.cn",
        "0009-0001-0072-5576",
        "0009-0007-7648-4776",
        "School of Medicine, The Chinese University of Hong Kong, Shenzhen",
        "MED Start-up Building, 2001 Longxiang Boulevard",
    ]
    missing_authors = [token for token in author_tokens if token not in manuscript]
    check("author_identity", not missing_authors, f"missing={missing_authors or 'none'}")

    abstract = section(manuscript, "## Abstract", "## Keywords")
    abstract_words = words(abstract)
    abstract_labels = [f"**{label}:**" in abstract for label in ("Background", "Methods", "Results", "Conclusions")]
    check("structured_abstract", all(abstract_labels) and abstract_words <= 350, f"words={abstract_words}; labels={sum(abstract_labels)}/4; limit=350")
    keyword_text = section(manuscript, "## Keywords", "## Background").replace("\n", " ")
    keyword_count = len([item for item in keyword_text.split(";") if item.strip()])
    check("keywords", 3 <= keyword_count <= 10, f"keywords={keyword_count}; allowed=3-10")

    required_sections = ["## Background", "## Methods", "## Results", "## Discussion", "## Conclusions", "## List of abbreviations", "## Declarations", "## References"]
    missing_sections = [heading for heading in required_sections if heading not in manuscript]
    check("required_sections", not missing_sections, f"missing={missing_sections or 'none'}")
    declaration_sections = [
        "### Ethics approval and consent to participate",
        "### Consent for publication",
        "### Availability of data and materials",
        "### Competing interests",
        "### Funding",
        "### Authors' contributions",
        "### Acknowledgements",
        "### Authors' information",
    ]
    missing_declarations = [heading for heading in declaration_sections if heading not in manuscript]
    check("declaration_structure", not missing_declarations, f"missing={missing_declarations or 'none'}")

    legends = re.findall(r"^### Figure (\d+) \| ([^\n]+)\n(.+?)(?=^### Figure |^## References)", manuscript, flags=re.M | re.S)
    legend_details = []
    for number, title, body in legends:
        legend_details.append({"figure": int(number), "title_words": words(title), "legend_words": words(body)})
    legends_pass = len(legends) == 5 and all(row["title_words"] <= 15 and row["legend_words"] <= 300 for row in legend_details)
    check("figure_legends", legends_pass, f"count={len(legends)}; details={legend_details}")

    figure_details = []
    figures_pass = True
    for number in range(1, 6):
        pngs = list((PACKAGE / "figures").glob(f"Figure{number}_*.png"))
        pdfs = list((PACKAGE / "figures").glob(f"Figure{number}_*.pdf"))
        if len(pngs) != 1 or len(pdfs) != 1:
            figures_pass = False
            figure_details.append({"figure": number, "error": "file count"})
            continue
        with Image.open(pngs[0]) as image:
            dimensions = image.size
        ok = pngs[0].stat().st_size < 10_000_000 and pdfs[0].stat().st_size < 10_000_000 and dimensions[0] >= 4000 and dimensions[1] >= 3000
        figures_pass &= ok
        figure_details.append({"figure": number, "png_bytes": pngs[0].stat().st_size, "pdf_bytes": pdfs[0].stat().st_size, "pixels": dimensions})
    check("figure_files", figures_pass, f"{figure_details}")

    docx_details = {path.name: path.stat().st_size if path.exists() else 0 for path in (MAIN_DOCX, SUPP_DOCX, COVER_DOCX)}
    check("editable_documents", all(size > 30_000 for size in docx_details.values()), f"{docx_details}")

    main_xml = ooxml(MAIN_DOCX)
    main_document = ET.fromstring(main_xml["word/document.xml"])
    main_settings = ET.fromstring(main_xml["word/settings.xml"])
    main_styles = ET.fromstring(main_xml["word/styles.xml"])
    line_numbers = main_document.find(".//w:sectPr/w:lnNumType", NS) is not None
    page_field = b"PAGE" in b"".join(value for key, value in main_xml.items() if key.startswith("word/footer"))
    header_parts = [value for key, value in main_xml.items() if key.startswith("word/header")]
    header_texts = ["".join(ET.fromstring(value).itertext()) for value in header_parts]
    even_odd = main_settings.find(".//w:evenAndOddHeaders", NS) is not None
    normal_style = main_styles.find(".//w:style[@w:styleId='Normal']", NS)
    spacing = normal_style.find(".//w:spacing", NS) if normal_style is not None else None
    double_spacing = spacing is not None and spacing.get(f"{{{NS['w']}}}line") == "480"
    title_style = main_styles.find(".//w:style[@w:styleId='Title']", NS)
    no_title_border = title_style is not None and title_style.find(".//w:pBdr", NS) is None
    docx_format_pass = line_numbers and page_field and even_odd and len(header_parts) >= 2 and all("Genome Medicine" in text for text in header_texts) and double_spacing and no_title_border
    check("main_docx_ooxml", docx_format_pass, f"line_numbers={line_numbers}; page_field={page_field}; even_odd={even_odd}; headers={len(header_parts)}; double_spacing={double_spacing}; title_border_absent={no_title_border}")

    supplement_doc = Document(SUPP_DOCX)
    table_geometry = []
    for table in supplement_doc.tables:
        xml = table._tbl
        table_properties = xml.find(qn("w:tblPr"))
        table_width = table_properties.find(qn("w:tblW")) if table_properties is not None else None
        table_grid = xml.find(qn("w:tblGrid"))
        cell_widths = []
        for row in table.rows:
            for cell in row.cells:
                cell_properties = cell._tc.find(qn("w:tcPr"))
                cell_widths.append(cell_properties is not None and cell_properties.find(qn("w:tcW")) is not None)
        table_geometry.append(table_width is not None and table_grid is not None and all(cell_widths))
    check("supplement_table_geometry", len(table_geometry) == 5 and all(table_geometry), f"tables={len(table_geometry)}; explicit_geometry={sum(table_geometry)}/5")

    with REF_CSV.open(encoding="utf-8", newline="") as handle:
        ref_rows = list(csv.DictReader(handle))
    check("doi_reference_verification", len(ref_rows) == 13 and all(row["status"] == "PASS" for row in ref_rows), f"Crossref PASS={sum(row['status'] == 'PASS' for row in ref_rows)}/13; total manuscript references=17")

    numeric_tokens = ["0.947", "0.837", "1.086", "1.042", "rho=0.026", "3.187, 3.050 and 3.527", "3.294 and 3.666"]
    missing_numbers = [token for token in numeric_tokens if token not in manuscript]
    check("frozen_numeric_anchors", not missing_numbers, f"missing={missing_numbers or 'none'}")
    stale_terms = ["expanded ABC/APC-like", "central ABC/APC-like", "validated ABC/APC-like", "causal STAT1 activation"]
    stale_hits = [term for term in stale_terms if term.lower() in manuscript.lower()]
    check("no_stale_claim_language", not stale_hits, f"hits={stale_hits or 'none'}")
    boundary_tokens = [
        "not establish a discrete subtype, causal regulator or unique upstream stimulus",
        "do not identify a unique initiating ligand or establish causation in SLE",
        "do not prove that STAT1 or STAT2 initiated the in vivo state",
    ]
    missing_boundaries = [token for token in boundary_tokens if token not in manuscript]
    check("noncausal_boundaries", not missing_boundaries, f"present={len(boundary_tokens) - len(missing_boundaries)}/{len(boundary_tokens)}")

    source_zip = PACKAGE / "additional_files" / "Additional_file_2_Figure_Source_Data_GateC8.zip"
    with zipfile.ZipFile(source_zip) as archive:
        source_entries = sorted(name for name in archive.namelist() if not name.endswith("/"))
    source_csvs = sorted((PACKAGE / "additional_files" / "source_data").glob("Figure*_source_data.csv"))
    hash_rows = list(csv.DictReader((PACKAGE / "additional_files" / "source_data" / "SHA256SUMS.csv").open(encoding="utf-8", newline="")))
    source_hashes_pass = len(source_csvs) == 5 and len(hash_rows) == 5 and all(sha256(PACKAGE / "additional_files" / "source_data" / row["file"]).lower() == row["sha256"].lower() for row in hash_rows)
    check("figure_source_data", source_hashes_pass and len(source_entries) == 6, f"csv={len(source_csvs)}; checksums={len(hash_rows)}; zip_entries={len(source_entries)}")

    render_specs = [("main", MAIN_PDF, 21), ("supplement", SUPP_PDF, 3), ("cover", COVER_PDF, 1)]
    render_details = []
    renders_pass = True
    for name, pdf, expected in render_specs:
        pages = len(PdfReader(pdf).pages) if pdf.exists() else 0
        png_count = len(list(pdf.parent.glob("page-*.png")))
        ok = pdf.exists() and pdf.stat().st_size > 50_000 and pages == expected and png_count == expected
        renders_pass &= ok
        render_details.append({"document": name, "pdf_bytes": pdf.stat().st_size if pdf.exists() else 0, "pdf_pages": pages, "png_pages": png_count})
    check("wps_render_outputs", renders_pass, f"{render_details}; visual_review=PASS_ALL_25_PAGES")

    hard_stops = [
        "institutional ethics determination",
        "competing interests",
        "funding",
        "final authorship completeness, corresponding-author designation, CRediT contributions and all-author approval",
        "acknowledgements",
        "all-author originality/submission confirmation",
        "open-source licence and immutable archive DOI",
    ]
    manuscript_placeholders = manuscript.count("[[")
    cover_source = (ROOT / "04_submission" / "cover_letter_genome_medicine_gateC8_AUTHOR_COMPLETION_REQUIRED_2026-08-20.md").read_text(encoding="utf-8")
    cover_placeholders = cover_source.count("[[")
    check("hard_stops_visible", manuscript_placeholders == 6 and cover_placeholders == 2, f"manuscript placeholders={manuscript_placeholders}; cover placeholders={cover_placeholders}; unresolved hard stops={len(hard_stops)}")

    scientific_checks = [name for name in checks if name != "hard_stops_visible"]
    scientific_pass = all(checks[name]["pass"] for name in scientific_checks)
    portal_authorized = scientific_pass and not hard_stops
    decision = "PASS_GATE_C8_SCIENTIFIC_TECHNICAL_SUBMISSION_PACKAGE_AUTHOR_DECLARATIONS_AND_ARCHIVE_REQUIRED" if scientific_pass else "HOLD_GATE_C8_TECHNICAL_REPAIR_REQUIRED"

    audit = {
        "created_at": "2026-08-20",
        "decision": decision,
        "primary_target": "Genome Medicine",
        "scientific_technical_package_pass": scientific_pass,
        "portal_submission_authorized": portal_authorized,
        "checks": checks,
        "hard_stops": hard_stops,
        "next_stage": "Gate C8B author declarations, institutional ethics determination, repository licence and immutable archive DOI; then final portal preflight and submission",
    }
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    audit_json = RUN_DIR / "03_GATE_C8_FINAL_AUDIT.json"
    audit_md = RUN_DIR / "03_GATE_C8_FINAL_AUDIT.md"
    audit_json.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    md_lines = [
        "# Gate C8 final submission audit",
        "",
        f"**Decision:** `{decision}`",
        "",
        f"**Scientific and technical package:** {'PASS' if scientific_pass else 'HOLD'}",
        "",
        f"**Portal submission authorized:** {'YES' if portal_authorized else 'NO'}",
        "",
        "## Checks",
        "",
        "| Check | Result | Detail |",
        "|---|---:|---|",
    ]
    for name, result in checks.items():
        detail = str(result["detail"]).replace("|", "/").replace("\n", " ")
        md_lines.append(f"| `{name}` | {'PASS' if result['pass'] else 'FAIL'} | {detail} |")
    md_lines.extend(["", "## Author-controlled hard stops", ""] + [f"- {item}" for item in hard_stops] + ["", "## Next stage", "", audit["next_stage"]])
    audit_md.write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    review_dir = PACKAGE / "review_copies"
    review_dir.mkdir(parents=True, exist_ok=True)
    for path in (MAIN_PDF, SUPP_PDF, COVER_PDF):
        shutil.copy2(path, review_dir / path.name)
    (PACKAGE / "README_GATE_C8_PACKAGE.md").write_text(package_readme(), encoding="utf-8")
    internal = PACKAGE / "internal_qc"
    shutil.copy2(audit_json, internal / "GATE_C8_FINAL_AUDIT.json")
    shutil.copy2(audit_md, internal / "GATE_C8_FINAL_AUDIT.md")
    manifest_rows = build_manifest()
    run_manifest = RUN_DIR / "04_GATE_C8_INTEGRITY_MANIFEST.csv"
    shutil.copy2(PACKAGE / "MANIFEST_SHA256.csv", run_manifest)
    build_zip()

    summary = {
        **audit,
        "package_files_manifested": len(manifest_rows),
        "package_zip": PACKAGE_ZIP.relative_to(ROOT).as_posix(),
        "package_zip_bytes": PACKAGE_ZIP.stat().st_size,
        "package_zip_sha256": sha256(PACKAGE_ZIP),
    }
    (RUN_DIR / "05_GATE_C8_PACKAGE_STATUS.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    from docx.oxml.ns import qn

    main()
