#!/usr/bin/env python3
"""Finalize the QiTeng reader-language submission candidate and integrity archive."""

from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import shutil
import zipfile

from docx import Document
from pypdf import PdfReader, PdfWriter


ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "phase17_v7/npj_sba_qiteng_round3/20260929_reader_language_hard_negative"
BASELINE = RUN / "input_r2/SLE_Bcell_Submission_Package"
CANDIDATE = RUN / "candidate/SLE_Bcell_Submission_Package"
DELIVERY = ROOT / "04_submission/20260929_qiteng_submission_candidate"
PACKAGE = DELIVERY / "SLE_Bcell_Submission_Package"
ZIP_PATH = DELIVERY / "SLE_Bcell_Submission_Package_2026-09-29.zip"
S10_PDF = RUN / "artwork/supplementary_figures/Supplementary_Figure_S10_overlap_depletion.pdf"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def extracted_text(path: Path) -> str:
    document = Document(path)
    lines: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text:
            lines.append(f"[{paragraph.style.name}] {paragraph.text}")
    for number, table in enumerate(document.tables, start=1):
        lines.append("")
        lines.append(f"[TABLE {number}]")
        for row in table.rows:
            lines.append(" | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells))
    return "\n".join(lines).rstrip() + "\n"


def merge_review_pdf() -> int:
    writer = PdfWriter()
    inputs = [PACKAGE / "Manuscript/Manuscript.pdf"] + [
        PACKAGE / f"Figures/Figure_{number}.pdf" for number in range(1, 6)
    ]
    page_count = 0
    for path in inputs:
        reader = PdfReader(path)
        page_count += len(reader.pages)
        for page in reader.pages:
            writer.add_page(page)
    output = PACKAGE / "Author_Resources/Manuscript_With_Figures.pdf"
    with output.open("wb") as handle:
        writer.write(handle)
    return page_count


def update_provenance() -> None:
    path = PACKAGE / "Author_Resources/Source_Provenance.json"
    records = json.loads(path.read_text(encoding="utf-8"))
    by_path = {record["package_path"]: record for record in records}
    updates = {
        "Figures/Figure_2.pdf": {
            "source_path": str((RUN / "artwork/figures/Figure2_sample_level_composition.pdf").relative_to(ROOT)).replace("\\", "/"),
            "parent_sha256": sha256(BASELINE / "Figures/Figure_2.pdf"),
            "revision_note": "Source-level reader-language rebuild: 'Primary >=50'; scientific source data are byte-identical to the parent package.",
        },
        "Figures/Figure_3.pdf": {
            "source_path": str((RUN / "artwork/figures/Figure3_gse174188_bconv_transcription.pdf").relative_to(ROOT)).replace("\\", "/"),
            "parent_sha256": sha256(BASELINE / "Figures/Figure_3.pdf"),
            "revision_note": "Source-level reader-language rebuild: prespecified-program terminology; scientific source data are byte-identical to the parent package.",
        },
        "Supplementary_Information/Supplementary_Information.docx": {
            "source_path": str((RUN / "candidate/SLE_Bcell_Submission_Package/Supplementary_Information/Supplementary_Information.docx").relative_to(ROOT)).replace("\\", "/"),
            "parent_sha256": sha256(BASELINE / "Supplementary_Information/Supplementary_Information.docx"),
            "revision_note": "Rebuilt from canonical Markdown with reader-facing terminology, retained S7 readability layout and source-redrawn S10.",
        },
        "Supplementary_Information/Supplementary_Information.pdf": {
            "source_path": str((RUN / "candidate/SLE_Bcell_Submission_Package/Supplementary_Information/Supplementary_Information.pdf").relative_to(ROOT)).replace("\\", "/"),
            "parent_sha256": sha256(BASELINE / "Supplementary_Information/Supplementary_Information.pdf"),
            "revision_note": "WPS export of the rebuilt Supplementary Information; 15 pages and cross-render checked with LibreOffice.",
        },
        "Author_Resources/Supplementary_Figures/Supplementary_Figure_S10.pdf": {
            "source_path": str(S10_PDF.relative_to(ROOT)).replace("\\", "/"),
            "parent_sha256": sha256(BASELINE / "Author_Resources/Supplementary_Figures/Supplementary_Figure_S10.pdf"),
            "revision_note": "Source-level terminology rebuild only; 36-row depletion source data are byte-identical to the frozen archive.",
        },
        "Author_Resources/Manuscript_With_Figures.pdf": {
            "source_path": "assembled from current Manuscript/Manuscript.pdf and Figures/Figure_1.pdf through Figure_5.pdf",
            "parent_sha256": sha256(BASELINE / "Author_Resources/Manuscript_With_Figures.pdf"),
            "revision_note": "Reviewer convenience PDF rebuilt from the current 31-page manuscript and five current one-page main figures.",
        },
    }
    for package_path, fields in updates.items():
        record = by_path[package_path]
        record.update(fields)
        record["sha256"] = sha256(PACKAGE / package_path)
        record["revision_date"] = "2026-09-29"
        record["derived_via"] = "QiTeng Academic Writing Skill v0.3.24.2 reader-language and hard-negative gate"
    path.write_text(json.dumps(records, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def write_readme() -> None:
    text = """# SLE B-cell remodeling submission files

Target: npj Systems Biology and Applications. Article type: Article.

Upload `Manuscript/Manuscript.docx` (or its PDF), the five `Figures/Figure_*.pdf` files, `Supplementary_Information/Supplementary_Information.pdf`, the three Supplementary Data ZIPs, and `Cover_Letter/Cover_Letter.docx` or PDF under the portal's corresponding categories. Select one format per document unless the portal requests both. The combined review PDF in `Author_Resources` is optional and must not be uploaded as a second manuscript.

`Author_Resources` contains author metadata, declarations, reporting information, separately exported supplementary figures and integrity/provenance files. These are preparation resources, not an additional research supplement. The ten supplementary figures are already embedded in the single Supplementary Information PDF.

Author order: Zhi Chen, then Teng Qi. Corresponding author: Teng Qi. The package retains the author-approved funding, competing-interest, ethics, originality and generative-AI statements. The official reporting forms and account-level portal fields must still be completed in the journal interface when requested; this package has not been submitted.

The DOI `10.5281/zenodo.22151739` identifies the published scientific archive. The current presentation-only refinements do not replace that scientific archive: Figure 2, Figure 3 and Supplementary Figure S10 were rebuilt from their source code solely to replace internal workflow labels with reader-facing `primary` or `prespecified` terminology. Their machine-readable scientific source data are unchanged.

The current journal guidance and Nature Portfolio editorial/AI policies were rechecked on 29 September 2026. The package is suitable as a scientific submission candidate; final portal opening, metadata entry, file-category assignment and proof inspection remain human-controlled steps.

Run `python Author_Resources/Verify_Package.py` from any directory to verify the extracted package. `SHA256SUMS.csv` lists every payload file; the enclosing ZIP checksum is reported in the accompanying action record.
"""
    (PACKAGE / "README.md").write_text(text, encoding="utf-8", newline="\n")


def write_manifest() -> int:
    manifest = PACKAGE / "SHA256SUMS.csv"
    files = sorted(
        path for path in PACKAGE.rglob("*") if path.is_file() and path != manifest
    )
    with manifest.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["relative_path", "bytes", "sha256"])
        writer.writeheader()
        for path in files:
            writer.writerow(
                {
                    "relative_path": path.relative_to(PACKAGE).as_posix(),
                    "bytes": path.stat().st_size,
                    "sha256": sha256(path),
                }
            )
    return len(files)


def write_zip() -> None:
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(PACKAGE.rglob("*")):
            if not path.is_file():
                continue
            name = f"{PACKAGE.name}/{path.relative_to(PACKAGE).as_posix()}"
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 29, 12, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            archive.writestr(info, path.read_bytes(), compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)


def changed_files() -> list[dict[str, object]]:
    paths = sorted(
        {path.relative_to(BASELINE).as_posix() for path in BASELINE.rglob("*") if path.is_file()}
        | {path.relative_to(PACKAGE).as_posix() for path in PACKAGE.rglob("*") if path.is_file()}
    )
    rows = []
    for relative in paths:
        before = BASELINE / relative
        after = PACKAGE / relative
        before_hash = sha256(before) if before.exists() else ""
        after_hash = sha256(after) if after.exists() else ""
        if before_hash != after_hash:
            rows.append(
                {
                    "relative_path": relative,
                    "change": "modified" if before.exists() and after.exists() else "added_or_removed",
                    "parent_sha256": before_hash,
                    "candidate_sha256": after_hash,
                }
            )
    with (RUN / "changed_files_final.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]) if rows else ["relative_path"])
        writer.writeheader()
        writer.writerows(rows)
    return rows


def main() -> None:
    if DELIVERY.exists():
        raise RuntimeError(f"Refusing to overwrite existing delivery directory: {DELIVERY}")
    if not CANDIDATE.exists() or not BASELINE.exists() or not S10_PDF.exists():
        raise FileNotFoundError("Required candidate, baseline or rebuilt S10 is missing")
    DELIVERY.mkdir(parents=True)
    shutil.copytree(CANDIDATE, PACKAGE)
    shutil.copyfile(S10_PDF, PACKAGE / "Author_Resources/Supplementary_Figures/Supplementary_Figure_S10.pdf")
    supplement_docx = PACKAGE / "Supplementary_Information/Supplementary_Information.docx"
    (PACKAGE / "Supplementary_Information/Supplementary_Information.extracted.txt").write_text(
        extracted_text(supplement_docx), encoding="utf-8", newline="\n"
    )
    review_pages = merge_review_pdf()
    if review_pages != 36:
        raise RuntimeError(f"Expected 36 review pages; found {review_pages}")
    write_readme()
    update_provenance()
    manifest_rows = write_manifest()
    write_zip()
    changes = changed_files()
    status = {
        "status": "FINALIZED_QITENG_SCIENTIFIC_SUBMISSION_CANDIDATE_AUDIT_REQUIRED",
        "package": str(PACKAGE.relative_to(ROOT)).replace("\\", "/"),
        "zip": str(ZIP_PATH.relative_to(ROOT)).replace("\\", "/"),
        "zip_bytes": ZIP_PATH.stat().st_size,
        "zip_sha256": sha256(ZIP_PATH),
        "manifest_payloads": manifest_rows,
        "review_pdf_pages": review_pages,
        "changed_files_vs_round2": len(changes),
    }
    (RUN / "finalize_status.json").write_text(
        json.dumps(status, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    print(json.dumps(status, indent=2))


if __name__ == "__main__":
    main()
