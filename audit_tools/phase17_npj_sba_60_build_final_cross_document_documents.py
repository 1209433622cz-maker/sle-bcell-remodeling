#!/usr/bin/env python3
"""Build manuscript and Supplementary documents for the maintenance freeze."""

from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

import phase17_c8s_04_build_documents as documents


ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "phase17_v7/npj_sba_scientific_presentation_maintenance_freeze/20260908_final_cross_document"
SOURCES = RUN / "sources"
FIGURES = RUN / "figures/figures"
DOCUMENTS = RUN / "documents"
MAIN_STEM = "Manuscript_Scientific_Presentation_Freeze"
SUPPLEMENT_STEM = "Supplementary_Information_Scientific_Presentation_Freeze"
TITLE = (
    "Disease-blind reconstruction distinguishes reproducible interferon remodeling from less stable "
    "B-cell state assignments in systemic lupus erythematosus"
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def extract_text(document: Document) -> str:
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def remove_table_spacers_before_manual_page_breaks(document: Document) -> int:
    removed = 0
    for paragraph in list(document.paragraphs):
        if not paragraph._p.xpath('.//w:br[@w:type="page"]'):
            continue
        previous = paragraph._p.getprevious()
        if previous is None or previous.tag != qn("w:p"):
            continue
        text = "".join(previous.itertext()).strip()
        has_break = bool(previous.xpath('.//w:br[@w:type="page"]'))
        has_drawing = bool(previous.xpath(".//w:drawing"))
        if not text and not has_break and not has_drawing:
            previous.getparent().remove(previous)
            removed += 1
    return removed


def remove_manual_page_break_before_heading(document: Document, heading: str) -> int:
    removed = 0
    for paragraph in document.paragraphs:
        if paragraph.text.strip() != heading:
            continue
        previous = paragraph._p.getprevious()
        if previous is None or previous.tag != qn("w:p"):
            continue
        if previous.xpath('.//w:br[@w:type="page"]'):
            previous.getparent().remove(previous)
            removed += 1
    return removed


def patch_main(path: Path) -> dict[str, int]:
    document = Document(path)
    compacted = 0
    for paragraph in document.paragraphs:
        if re.fullmatch(r"Figure [1-5] \| .+", paragraph.text.strip()):
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(2)
            compacted += 1
    document.core_properties.author = "Zhi Chen; Teng Qi"
    document.core_properties.title = TITLE
    document.core_properties.subject = "Scientific-presentation final cross-document freeze"
    document.core_properties.comments = (
        "Discussion landing synchronized with identity, transfer and mechanistic limits; "
        "no estimate, figure or Source Data change."
    )
    document.save(path)
    return {"main_legend_headings_compacted": compacted}


def patch_supplement(path: Path) -> dict[str, int]:
    document = Document(path)
    removed_spacers = remove_table_spacers_before_manual_page_breaks(document)
    removed_s8 = remove_manual_page_break_before_heading(
        document, "Supplementary Table S8 | Full statistical-results archive map"
    )
    removed_s1 = remove_manual_page_break_before_heading(
        document, "Supplementary Figure S1 | Source integrity and hard-quality-control diagnostics"
    )
    removed_s2 = remove_manual_page_break_before_heading(
        document, "Supplementary Figure S2 | Representation and bridge diagnostics"
    )
    s5_tables = [
        table
        for table in document.tables
        if [cell.text.strip() for cell in table.rows[0].cells]
        == ["Figure", "Evidence basis", "Machine-readable source"]
    ]
    if len(s5_tables) != 1:
        raise RuntimeError(f"Expected one Supplementary Table S5, found {len(s5_tables)}")
    s5 = s5_tables[0]
    widths_mm = (28.0, 92.0, 45.0)
    for column, width in zip(s5.columns, widths_mm, strict=True):
        column.width = Mm(width)
        for cell in column.cells:
            cell.width = Mm(width)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
    document.core_properties.author = "Zhi Chen; Teng Qi"
    document.core_properties.title = "Supplementary information"
    document.core_properties.subject = "Scientific-presentation final cross-document freeze"
    document.core_properties.comments = (
        "Selected figure source-data map synchronized with Figure 1d, 4d and 5d; "
        "no figure or Source Data change."
    )
    document.save(path)
    return {
        "redundant_spacers_removed": removed_spacers,
        "renderer_divergent_s8_breaks_removed": removed_s8,
        "s1_breaks_removed": removed_s1,
        "s2_breaks_removed": removed_s2,
        "s5_rows_compacted": len(s5.rows),
    }


def figure_alt_titles(document: Document) -> list[str | None]:
    return [shape._inline.docPr.get("title") for shape in document.inline_shapes]


def main() -> None:
    integration = json.loads((RUN / "06_FINAL_CROSS_DOCUMENT_INTEGRATION_STATUS.json").read_text(encoding="utf-8"))
    if integration["failed_checks"]:
        raise RuntimeError("Cross-document source integration contains failed checks")

    DOCUMENTS.mkdir(parents=True, exist_ok=True)
    main_source = SOURCES / "Manuscript_scientific_presentation_freeze.md"
    supplement_source = SOURCES / "Supplementary_Information_scientific_presentation_freeze.md"
    main_output = DOCUMENTS / f"{MAIN_STEM}.docx"
    supplement_output = DOCUMENTS / f"{SUPPLEMENT_STEM}.docx"

    build_results = {
        "manuscript": documents.markdown_to_docx(
            main_source,
            main_output,
            body_size=12,
            double_space=True,
            line_numbers=True,
            running_header="npj Systems Biology and Applications | Article",
            title_override=TITLE,
        ),
        "supplement": documents.markdown_to_docx(
            supplement_source,
            supplement_output,
            body_size=10.5,
            double_space=False,
            line_numbers=False,
            running_header="Supplementary information",
            title_override="Supplementary information",
            supplementary_figure_dirs=[FIGURES],
            page_break_before_headings=set(),
        ),
    }
    repairs = {"manuscript": patch_main(main_output), "supplement": patch_supplement(supplement_output)}

    manuscript = Document(main_output)
    supplement = Document(supplement_output)
    manuscript_text = extract_text(manuscript)
    supplement_text = extract_text(supplement)
    source_text = main_source.read_text(encoding="utf-8")
    reference_numbers = [
        int(value)
        for value in re.findall(
            r"(?m)^(\d+)\. ",
            source_text.split("## References\n", 1)[1].split("## Figure legends\n", 1)[0],
        )
    ]
    abstract_match = re.search(r"Abstract\n(.+?)\nIntroduction", manuscript_text, flags=re.DOTALL)
    if not abstract_match:
        raise RuntimeError("Could not identify manuscript abstract")
    abstract_words = re.findall(r"\b[\w'-]+\b", abstract_match.group(1))
    expected_alt_titles = [f"Supplementary Figure S{number}" for number in range(1, 11)]

    checks = {
        "manuscript_title_exact": TITLE in manuscript_text,
        "abstract_145_words": len(abstract_words) == 145,
        "references_1_to_33": reference_numbers == list(range(1, 34)),
        "main_has_no_inline_figures": len(manuscript.inline_shapes) == 0,
        "five_main_legend_headings_compacted": repairs["manuscript"]["main_legend_headings_compacted"] == 5,
        "discussion_landing_exact": "within explicit identity, transfer and mechanistic limits." in manuscript_text,
        "supplement_has_ten_figures": len(supplement.inline_shapes) == 10,
        "supplement_has_ten_grid_objects": len(supplement.tables) == 10,
        "supplement_alt_titles_s1_to_s10": figure_alt_titles(supplement) == expected_alt_titles,
        "supplement_heading_sequence_s1_to_s10": [
            int(value) for value in re.findall(r"Supplementary Figure S(10|[1-9]) \|", supplement_text)
        ] == list(range(1, 11)),
        "s5_figure1_owner_updated": "Disease-blind identity stability, two-compartment adjudication and end-to-end B_ASC boundary" in supplement_text,
        "s5_figure4_owner_updated": "Source-label-defined GSE135779 replication, gene-level coherence and required calibration boundary" in supplement_text,
        "s5_figure5_owner_updated": "Regulatory convergence, IFN-overlap-depletion ceiling and orthogonal response evidence" in supplement_text,
        "s5_nine_rows_locally_compacted": repairs["supplement"]["s5_rows_compacted"] == 9,
        "redundant_s2_break_removed": repairs["supplement"]["s2_breaks_removed"] == 1,
    }
    failed = [name for name, passed in checks.items() if not passed]
    files = {
        path.name: {
            "path": path.relative_to(ROOT).as_posix(),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        }
        for path in (main_output, supplement_output)
    }
    status = {
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "status": "PASS_FINAL_CROSS_DOCUMENT_DOCUMENTS_BUILT_RENDER_REQUIRED" if not failed else "FAIL_FINAL_CROSS_DOCUMENT_DOCUMENT_BUILD",
        "build_results": build_results,
        "layout_repairs": repairs,
        "checks": checks,
        "failed_checks": failed,
        "files": files,
        "scientific_estimates_changed": False,
        "figures_redrawn": False,
        "source_data_values_changed": False,
    }
    (RUN / "07_DOCUMENT_BUILD_STATUS.json").write_text(
        json.dumps(status, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    print(json.dumps(status, indent=2))
    if failed:
        raise RuntimeError(f"Document build checks failed: {failed}")


if __name__ == "__main__":
    main()
