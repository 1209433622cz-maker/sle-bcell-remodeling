"""Rebuild the S7 layout from frozen Markdown without editing scientific content."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import phase17_npj_sba_60_build_final_cross_document_documents as parent

ROOT = parent.ROOT
RUN = ROOT / 'phase17_v7/npj_sba_s7_readability/20260909_source_rebuild'
STEM = 'Supplementary_Information_S7_readability'

def build():
    out = RUN / 'documents' / (STEM + '.docx')
    out.parent.mkdir(parents=True, exist_ok=True)
    parent.documents.markdown_to_docx(
        parent.SOURCES / 'Supplementary_Information_scientific_presentation_freeze.md',
        out, body_size=10.5, double_space=False, line_numbers=False,
        running_header='Supplementary information', title_override='Supplementary information',
        supplementary_figure_dirs=[parent.FIGURES], page_break_before_headings=set())
    parent.patch_supplement(out)
    doc = Document(out)
    table = next(t for t in doc.tables if t.cell(0, 0).text == 'Result family')
    widths = [165.1 * ratio / 130 for ratio in (19, 27, 24, 19, 23, 18)]
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Mm(width)
        for cell in column.cells:
            cell.width = Mm(width)
            pr = cell._tc.get_or_add_tcPr()
            margin = pr.find(qn('w:tcMar'))
            if margin is None:
                margin = OxmlElement('w:tcMar')
                pr.append(margin)
            for side in ('left', 'right', 'start', 'end'):
                element = OxmlElement('w:' + side)
                element.set(qn('w:w'), '50')
                element.set(qn('w:type'), 'dxa')
                margin.append(element)
    # The shorter S7 frees page space; keep S2 on its own existing figure page.
    for p in doc.paragraphs:
        if p.text.startswith('Supplementary Figure S2 |'):
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(0)
    original = Document(parent.DOCUMENTS / (parent.SUPPLEMENT_STEM + '.docx'))
    assert parent.extract_text(doc) == parent.extract_text(original)
    assert [s._inline.docPr.get('title') for s in doc.inline_shapes] == parent.figure_alt_titles(original)
    doc.save(out)
    (RUN / 'build.json').write_text(json.dumps({
        's7_column_widths_mm': widths, 'horizontal_padding_twips': 50,
        'text_identical': True, 'source': str(parent.SOURCES),
        'font_size_changed': False, 's2_page_break_before': True,
        'docx_sha256': parent.sha256(out)}, indent=2))
    print(out)

if __name__ == '__main__':
    build()
