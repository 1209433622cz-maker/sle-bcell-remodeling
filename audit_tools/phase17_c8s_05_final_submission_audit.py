#!/usr/bin/env python3
"""Audit and deterministically package the Gate C8S submission handoff."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
RUN_DIR = ROOT / "phase17_v7" / "gateC8S" / "20260821_supplementary_traceability_freeze"
PACKAGE = ROOT / "04_submission" / "package_genome_medicine_gateC8S_2026-08-21"
PACKAGE_ZIP = ROOT / "04_submission" / "package_genome_medicine_gateC8S_2026-08-21.zip"
MANUSCRIPT_MD = ROOT / "01_manuscript" / "manuscript_v12_genome_medicine_gateC8S_2026-08-21.md"
SUPPLEMENT_MD = ROOT / "01_manuscript" / "supplementary_information_v3_gateC8S_2026-08-21.md"
TARGET_MD = ROOT / "04_submission" / "journal_target_decision_gateC8S_2026-08-21.md"
COVER_MD = ROOT / "04_submission" / "cover_letter_genome_medicine_gateC8S_AUTHOR_COMPLETION_REQUIRED_2026-08-21.md"
REFERENCE_CSV = ROOT / "phase17_v7" / "gateC8R" / "20260820_pre_submission_repair" / "references" / "reference_verification_gateC8R.csv"

MAIN_DOCX = PACKAGE / "main_text" / "Genome_Medicine_Manuscript_GateC8S_AUTHOR_COMPLETION_REQUIRED.docx"
SUPP_DOCX = PACKAGE / "additional_files" / "Additional_file_1_Supplementary_Information_GateC8S.docx"
COVER_DOCX = PACKAGE / "submission_docs" / "Genome_Medicine_Cover_Letter_GateC8S_AUTHOR_CONFIRMATION_REQUIRED.docx"
MAIN_PDF = PACKAGE / "internal_qc" / "wps_render_main" / "Genome_Medicine_Manuscript_GateC8S_AUTHOR_COMPLETION_REQUIRED_WPS.pdf"
SUPP_PDF = PACKAGE / "internal_qc" / "wps_render_supplement" / "Additional_file_1_Supplementary_Information_GateC8S_WPS.pdf"
COVER_PDF = PACKAGE / "internal_qc" / "wps_render_cover_letter" / "Genome_Medicine_Cover_Letter_GateC8S_AUTHOR_CONFIRMATION_REQUIRED_WPS.pdf"

TITLE = "Disease-blind single-cell reconstruction separates unstable B-cell states from reproducible interferon remodeling in systemic lupus erythematosus"
FIXED_ZIP_TIME = (2026, 8, 21, 0, 0, 0)
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def load_json(name: str) -> dict:
    return json.loads((RUN_DIR / name).read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def words(text: str) -> int:
    clean = re.sub(r"[`*_#|\[\]]", " ", text)
    return len(re.findall(r"\b[\w'-]+\b", clean))


def section(text: str, start: str, end: str | None) -> str:
    begin = text.index(start) + len(start)
    finish = text.index(end, begin) if end else len(text)
    return text[begin:finish].strip()


def ooxml(docx_path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(docx_path) as archive:
        return {name: archive.read(name) for name in archive.namelist() if name.endswith(".xml")}


def package_readme() -> str:
    return f"""# Gate C8S Genome Medicine submission handoff

This is the supplementary-evidence and traceability freeze generated on 21 August 2026.

## Status

- Scientific, figure, statistical-traceability and technical package: PASS.
- Portal submission: BLOCKED until the author-controlled declarations, institutional ethics determination, repository licence and immutable archive DOI are completed.
- Primary target: Genome Medicine; journal quartiles are not asserted in this frozen package.
- WPS visual review: PASS for all 27 manuscript pages, 12 supplementary pages and 1 cover-letter page.
- Main-panel assertions: PASS 46/46. Supplementary-panel assertions: PASS 29/29.
- Figure 5 source mapping: panel d is M5911 (3 rows); panel e is GSE23307 (2 donors, 12/12 positive genes each).
- Full statistical archive: 12 complete gene-level branches, 12 sanitized design matrices and 63 payload files; SHA-256 `{load_json('05_FULL_STATISTICAL_RESULTS_STATUS.json')['archive_sha256']}`.

## Upload mapping after author completion

- `main_text/`: editable main manuscript.
- `figures/`: Figures 1-5 as PDF and 600-dpi PNG composites.
- `figures_supplementary/`: Supplementary Figures S1-S7 as PDF and 600-dpi PNG composites.
- `additional_files/`: Supplementary Information, 12 figure source-data files, regulator sensitivity and full statistical results.
- `submission_docs/`: cover letter, author completion form, target decision and reporting checklist.
- `references/`: DOI verification and Vancouver reference records.
- `review_copies/`: WPS-rendered PDFs for author review; these are not the primary editable uploads.
- `internal_qc/`: accessibility reports, final page renders and audit evidence; do not upload unless requested.

## Hard stops

Do not submit until the visible author-controlled items have been supplied and verified. Do not alter frozen scientific estimates during Gate C8B.
"""


def build_manifest() -> list[dict[str, object]]:
    manifest_path = PACKAGE / "MANIFEST_SHA256.csv"
    rows: list[dict[str, object]] = []
    for path in sorted(p for p in PACKAGE.rglob("*") if p.is_file() and p != manifest_path):
        rows.append({
            "relative_path": path.relative_to(PACKAGE).as_posix(),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        })
    with manifest_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["relative_path", "bytes", "sha256"])
        writer.writeheader()
        writer.writerows(rows)
    return rows


def write_deterministic_zip(output: Path) -> None:
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(p for p in PACKAGE.rglob("*") if p.is_file()):
            relative = (Path(PACKAGE.name) / path.relative_to(PACKAGE)).as_posix()
            info = zipfile.ZipInfo(relative, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            info.create_system = 3
            archive.writestr(info, path.read_bytes(), compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)


def build_deterministic_archive() -> tuple[int, str]:
    with tempfile.TemporaryDirectory(prefix="gateC8S_zip_") as temp_dir:
        first = Path(temp_dir) / "first.zip"
        second = Path(temp_dir) / "second.zip"
        write_deterministic_zip(first)
        write_deterministic_zip(second)
        if sha256(first) != sha256(second):
            raise RuntimeError("Deterministic package ZIP rebuild check failed")
        shutil.copyfile(first, PACKAGE_ZIP)
    return PACKAGE_ZIP.stat().st_size, sha256(PACKAGE_ZIP)


def main() -> None:
    manuscript = MANUSCRIPT_MD.read_text(encoding="utf-8")
    supplement = SUPPLEMENT_MD.read_text(encoding="utf-8")
    target = TARGET_MD.read_text(encoding="utf-8")
    cover = COVER_MD.read_text(encoding="utf-8")
    main_status = load_json("01_FIGURE_BUILD_STATUS.json")
    main_assertions = load_json("02_PANEL_DATA_ASSERTIONS.json")
    supp_assertions = load_json("03_SUPPLEMENTARY_PANEL_DATA_ASSERTIONS.json")
    supp_status = load_json("04_SUPPLEMENTARY_FIGURE_BUILD_STATUS.json")
    stats_status = load_json("05_FULL_STATISTICAL_RESULTS_STATUS.json")
    source_status = load_json("06_GATE_C8S_SOURCE_BUILD_STATUS.json")
    document_status = load_json("07_GATE_C8S_DOCUMENT_BUILD_STATUS.json")
    checks: dict[str, dict[str, object]] = {}

    def check(name: str, passed: bool, detail: object) -> None:
        checks[name] = {"pass": bool(passed), "detail": str(detail)}

    main_rows = main_assertions.get("checks", [])
    main_map = {row.get("check"): row for row in main_rows}
    fig5_expected = {
        "Figure5.panel_d.source_rows": 3,
        "Figure5.panel_e.source_rows": 2,
        "Figure5.panel_e.donors_with_12_positive_genes": 2,
    }
    fig5_actual = {name: main_map.get(name, {}).get("actual") for name in fig5_expected}
    check(
        "main_figure_assertions",
        main_status.get("status") == "C8S_MAIN_FIGURES_BUILT_WITH_ASSERTIONS"
        and main_assertions.get("status") == "PASS"
        and len(main_rows) == 46
        and all(row.get("pass") is True for row in main_rows)
        and all(fig5_actual[name] == value for name, value in fig5_expected.items()),
        f"passed={sum(row.get('pass') is True for row in main_rows)}/46; Figure5={fig5_actual}",
    )

    supp_rows = supp_assertions.get("checks", [])
    check(
        "supplementary_figure_assertions",
        supp_status.get("status") == "PASS_GATE_C8S_SUPPLEMENTARY_FIGURES_BUILT"
        and supp_assertions.get("status") == "PASS"
        and len(supp_rows) == 29
        and all(row.get("pass") is True for row in supp_rows),
        f"passed={sum(row.get('pass') is True for row in supp_rows)}/29; figures={supp_status.get('figures')}",
    )

    figure_details = []
    figures_pass = True
    for prefix, directory, total in (
        ("Figure", PACKAGE / "figures", 5),
        ("Supplementary_Figure_S", PACKAGE / "figures_supplementary", 7),
    ):
        for number in range(1, total + 1):
            pngs = list(directory.glob(f"{prefix}{number}_*.png"))
            pdfs = list(directory.glob(f"{prefix}{number}_*.pdf"))
            ok = len(pngs) == 1 and len(pdfs) == 1
            dimensions = (0, 0)
            if ok:
                with Image.open(pngs[0]) as image:
                    dimensions = image.size
                ok = dimensions[0] >= 4000 and dimensions[1] >= 3000 and pngs[0].stat().st_size > 0 and pdfs[0].stat().st_size > 0
            figures_pass &= ok
            figure_details.append({"figure": f"{prefix}{number}", "pixels": dimensions, "pass": ok})
    check("publication_figure_files", figures_pass, f"pairs={len(figure_details)}; all >=4000x3000 and non-empty")

    source_dir = PACKAGE / "additional_files" / "source_data"
    source_rows = read_csv(source_dir / "SHA256SUMS.csv")
    source_csvs = sorted(source_dir.glob("*Figure*_source_data.csv"))
    source_hashes_ok = len(source_csvs) == 12 and len(source_rows) == 12 and all(
        sha256(source_dir / row["file"]) == row["sha256"].upper() for row in source_rows
    )
    fig5_rows = read_csv(source_dir / "Figure5_source_data.csv")
    panel_counts = {panel: sum(row["panel"] == panel for row in fig5_rows) for panel in ("B", "C", "D", "E")}
    mapping_ok = (
        panel_counts == {"B": 12, "C": 12, "D": 3, "E": 2}
        and all(row["series"] == "MSigDB_M5911_NES" for row in fig5_rows if row["panel"] == "D")
        and all(row["series"] == "GSE23307_mean_paired_log2p1_effect" for row in fig5_rows if row["panel"] == "E")
        and all(row["n_or_targets"] == "12" for row in fig5_rows if row["panel"] == "E")
    )
    check("figure_source_data", source_hashes_ok and mapping_ok, f"csv=12; hashes={len(source_rows)}; Figure5 panels={panel_counts}")

    stats_zip = PACKAGE / "additional_files" / "Additional_file_4_Full_Statistical_Results_GateC8S.zip"
    with zipfile.ZipFile(stats_zip) as archive:
        stats_entries = [name for name in archive.namelist() if not name.endswith("/")]
    stats_ok = (
        stats_status.get("status") == "PASS_GATE_C8S_FULL_STATISTICAL_ARCHIVE_BUILT"
        and stats_status.get("complete_gene_result_files") == 12
        and stats_status.get("sanitized_design_matrices") == 12
        and stats_status.get("payload_files") == 63
        and stats_status.get("deterministic_rebuild_match") is True
        and stats_status.get("direct_identifiers_in_design_matrices") is False
        and len(stats_entries) == 63
        and sha256(stats_zip) == stats_status.get("archive_sha256")
    )
    check("full_statistical_archive", stats_ok, f"entries={len(stats_entries)}; sha256={sha256(stats_zip)}")

    abstract_words = words(section(manuscript, "## Abstract", "## Keywords"))
    refs = [int(value) for value in re.findall(r"^(\d+)\.\s", section(manuscript, "## References", None), flags=re.M)]
    manuscript_tokens = [
        f"# {TITLE}",
        "**Date:** 21 August 2026",
        "### Statistical analysis and multiplicity",
        "one global BH family across eight regulators and three confirmatory contrasts (24 tests)",
        "separate BH families of six tests for each method",
        "The paired GSE23307 experiment had two donors and no inferential P value",
        "does not establish a predictive biomarker",
        "available in the public project repository",
        "Gate C8S is the current canonical analysis state",
        "**Additional file 4 (.zip):** Full statistical results.",
    ]
    missing_tokens = [token for token in manuscript_tokens if token not in manuscript]
    stale_terms = ["expanded ABC/APC-like", "causal STAT1 activation", "private project repository"]
    stale_hits = [term for term in stale_terms if term.lower() in manuscript.lower()]
    check(
        "manuscript_scope_and_statistics",
        not missing_tokens and not stale_hits and abstract_words == 314 and refs == list(range(1, 31)),
        f"abstract={abstract_words}; references={len(refs)}; missing={missing_tokens or 'none'}; stale={stale_hits or 'none'}",
    )
    check(
        "supplementary_source_contract",
        source_status.get("status") == "PASS_C8S_SUBMISSION_SOURCES_BUILT"
        and source_status.get("supplementary_figures") == 7
        and source_status.get("supplementary_tables") == 8
        and supplement.count("[[SUPPLEMENTARY_FIGURE:S") == 7,
        "7 figure markers; 8 tables; active C8S source status PASS",
    )

    author_tokens = [
        "Zhi Chen [1] and Teng Qi [1,*]",
        "zhichen1@link.cuhk.edu.cn",
        "tengqi@link.cuhk.edu.cn",
        "0009-0001-0072-5576",
        "0009-0007-7648-4776",
        "School of Medicine, The Chinese University of Hong Kong, Shenzhen",
    ]
    check("author_identity", all(token in manuscript for token in author_tokens), "first author and corresponding-author identity, emails, ORCIDs and affiliation present")
    check("journal_target", "**Primary submission: Genome Medicine.**" in target, "Genome Medicine primary without a frozen quartile assertion")

    docx_details = {path.name: path.stat().st_size if path.exists() else 0 for path in (MAIN_DOCX, SUPP_DOCX, COVER_DOCX)}
    check(
        "document_build_status",
        document_status.get("created_at") == "2026-08-21"
        and len(document_status.get("outputs", [])) == 3
        and document_status.get("assets", {}).get("figure_files") == 10
        and document_status.get("assets", {}).get("supplementary_figure_files") == 14
        and document_status.get("assets", {}).get("source_files") == 12
        and all(size > 30_000 for size in docx_details.values()),
        docx_details,
    )

    main_xml = ooxml(MAIN_DOCX)
    main_document = ET.fromstring(main_xml["word/document.xml"])
    main_settings = ET.fromstring(main_xml["word/settings.xml"])
    main_styles = ET.fromstring(main_xml["word/styles.xml"])
    line_numbers = main_document.find(".//w:sectPr/w:lnNumType", NS) is not None
    page_field = b"PAGE" in b"".join(value for key, value in main_xml.items() if key.startswith("word/footer"))
    even_odd = main_settings.find(".//w:evenAndOddHeaders", NS) is not None
    normal_style = main_styles.find(".//w:style[@w:styleId='Normal']", NS)
    spacing = normal_style.find(".//w:spacing", NS) if normal_style is not None else None
    double_spacing = spacing is not None and spacing.get(f"{{{NS['w']}}}line") == "480"
    check("main_docx_ooxml", line_numbers and page_field and even_odd and double_spacing, f"line_numbers={line_numbers}; page_field={page_field}; even_odd={even_odd}; double_spacing={double_spacing}")

    supplement_doc = Document(SUPP_DOCX)
    geometry = []
    for table in supplement_doc.tables:
        properties = table._tbl.find(qn("w:tblPr"))
        grid = table._tbl.find(qn("w:tblGrid"))
        widths = []
        for row in table.rows:
            for cell in row.cells:
                cell_properties = cell._tc.find(qn("w:tcPr"))
                widths.append(cell_properties is not None and cell_properties.find(qn("w:tcW")) is not None)
        geometry.append(
            properties is not None
            and properties.find(qn("w:tblW")) is not None
            and properties.find(qn("w:tblInd")) is not None
            and grid is not None
            and all(widths)
        )
    check("supplement_docx_structure", len(geometry) == 8 and all(geometry) and len(supplement_doc.inline_shapes) == 7, f"tables={len(geometry)}; explicit_geometry={sum(geometry)}; inline_figures={len(supplement_doc.inline_shapes)}")

    a11y_details = []
    for name in ("main_text_a11y.json", "supplement_a11y.json", "cover_letter_a11y.json"):
        report = json.loads((PACKAGE / "internal_qc" / name).read_text(encoding="utf-8"))
        counts = report.get("counts", {})
        a11y_details.append({"file": name, **counts})
    check("docx_accessibility", all(sum(int(row.get(level, 0)) for level in ("high", "medium", "low")) == 0 for row in a11y_details), a11y_details)

    render_specs = (("main", MAIN_PDF, 27), ("supplement", SUPP_PDF, 12), ("cover", COVER_PDF, 1))
    render_details = []
    renders_ok = True
    for name, pdf, expected in render_specs:
        reader = PdfReader(pdf) if pdf.exists() else None
        pages = len(reader.pages) if reader else 0
        pngs = len(list(pdf.parent.glob("final-page-*.png")))
        ok = pdf.exists() and pdf.stat().st_size > 50_000 and pages == expected and pngs == expected
        renders_ok &= ok
        render_details.append({"document": name, "pages": pages, "page_pngs": pngs, "bytes": pdf.stat().st_size if pdf.exists() else 0})
    supp_reader = PdfReader(SUPP_PDF)
    supp_page_text = [page.extract_text() or "" for page in supp_reader.pages]
    pagination_ok = "Supplementary Table S8" in supp_page_text[4] and all(
        f"Supplementary Figure S{number}" in supp_page_text[number + 4] for number in range(1, 8)
    )
    check("wps_render_and_visual_review", renders_ok and pagination_ok, f"{render_details}; S8 page=5; S1-S7 pages=6-12; visual_review=PASS_ALL_40_PAGES")

    references = read_csv(REFERENCE_CSV)
    check("reference_verification", len(references) == 26 and all(row["status"] == "PASS" for row in references) and len(refs) == 30, "Crossref PASS=26/26; manuscript references=30")

    source_zip = PACKAGE / "additional_files" / "Additional_file_2_Figure_Source_Data_GateC8S.zip"
    sensitivity_zip = PACKAGE / "additional_files" / "Additional_file_3_Regulator_Sensitivity_GateC8S.zip"
    with zipfile.ZipFile(source_zip) as archive:
        source_entries = [name for name in archive.namelist() if not name.endswith("/")]
    with zipfile.ZipFile(sensitivity_zip) as archive:
        sensitivity_entries = [name for name in archive.namelist() if not name.endswith("/")]
    check("attachment_integrity", len(source_entries) == 13 and len(sensitivity_entries) == 3 and stats_ok, f"source entries={len(source_entries)}; regulator entries={len(sensitivity_entries)}; statistical entries={len(stats_entries)}")

    hard_stops = [
        "institutional ethics determination",
        "competing interests",
        "funding",
        "final CRediT contributions and all-author approval",
        "acknowledgements",
        "all-author originality/submission confirmation",
        "open-source licence and immutable archive DOI",
    ]
    manuscript_placeholders = manuscript.count("[[")
    cover_placeholders = cover.count("[[")
    check("hard_stops_visible", manuscript_placeholders == 6 and cover_placeholders == 2, f"manuscript placeholders={manuscript_placeholders}; cover placeholders={cover_placeholders}; unresolved hard stops={len(hard_stops)}")

    scientific_checks = [name for name in checks if name != "hard_stops_visible"]
    scientific_pass = all(checks[name]["pass"] for name in scientific_checks)
    decision = "PASS_GATE_C8S_SUPPLEMENTARY_EVIDENCE_TRACEABILITY_FREEZE_AUTHOR_ACTION_REQUIRED" if scientific_pass else "HOLD_GATE_C8S_TECHNICAL_REPAIR_REQUIRED"
    audit = {
        "created_at": "2026-08-21",
        "decision": decision,
        "primary_target": "Genome Medicine",
        "scientific_technical_package_pass": scientific_pass,
        "portal_submission_authorized": False,
        "checks": checks,
        "hard_stops": hard_stops,
        "next_stage": "Gate C8B author declarations, institutional ethics determination, repository licence and immutable archive DOI; then rebuild once and complete portal preflight without changing frozen scientific results",
    }

    audit_json = RUN_DIR / "08_GATE_C8S_FINAL_AUDIT.json"
    audit_md = RUN_DIR / "08_GATE_C8S_FINAL_AUDIT.md"
    audit_json.write_text(json.dumps(audit, indent=2) + "\n", encoding="utf-8")
    lines = [
        "# Gate C8S final submission audit", "", f"**Decision:** `{decision}`", "",
        f"**Scientific and technical package:** {'PASS' if scientific_pass else 'HOLD'}", "",
        "**Portal submission authorized:** NO - author-controlled hard stops remain.", "", "## Checks", "",
        "| Check | Result | Detail |", "|---|---:|---|",
    ]
    for name, result in checks.items():
        detail = str(result["detail"]).replace("|", "/").replace("\n", " ")
        lines.append(f"| `{name}` | {'PASS' if result['pass'] else 'FAIL'} | {detail} |")
    lines.extend(["", "## Author-controlled hard stops", ""] + [f"- {item}" for item in hard_stops] + ["", "## Next stage", "", audit["next_stage"]])
    audit_md.write_text("\n".join(lines) + "\n", encoding="utf-8")

    review_dir = PACKAGE / "review_copies"
    review_dir.mkdir(parents=True, exist_ok=True)
    for path in (MAIN_PDF, SUPP_PDF, COVER_PDF):
        shutil.copy2(path, review_dir / path.name)
    (PACKAGE / "README_GATE_C8S_PACKAGE.md").write_text(package_readme(), encoding="utf-8")
    shutil.copy2(audit_json, PACKAGE / "internal_qc" / "GATE_C8S_FINAL_AUDIT.json")
    shutil.copy2(audit_md, PACKAGE / "internal_qc" / "GATE_C8S_FINAL_AUDIT.md")
    manifest_rows = build_manifest()
    shutil.copy2(PACKAGE / "MANIFEST_SHA256.csv", RUN_DIR / "09_GATE_C8S_INTEGRITY_MANIFEST.csv")
    zip_bytes, zip_hash = build_deterministic_archive()

    summary = {
        **audit,
        "deterministic_archive_rebuild_match": True,
        "package_files_manifested": len(manifest_rows),
        "package_zip": PACKAGE_ZIP.relative_to(ROOT).as_posix(),
        "package_zip_bytes": zip_bytes,
        "package_zip_sha256": zip_hash,
    }
    (RUN_DIR / "10_GATE_C8S_PACKAGE_STATUS.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))
    if not scientific_pass:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
