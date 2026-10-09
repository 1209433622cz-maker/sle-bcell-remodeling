"""Assemble the current initial-submission files with exact scientific provenance."""
import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import shutil
import zipfile
from docx import Document
from docx.shared import Pt
from pypdf import PdfReader, PdfWriter
import phase17_c8s_04_build_documents as docs

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'04_submission/20260909_current_submission'
PKG=OUT/'SLE_Bcell_Submission_Package'
BASE=ROOT/'phase17_v7/npj_sba_scientific_presentation_maintenance_freeze/20260908_final_cross_document'
S7=ROOT/'phase17_v7/npj_sba_s7_readability/20260909_source_rebuild'
OLD=ROOT/'04_submission/current_submission_package/SLE_Bcell_Submission_Package'
TITLE='Disease-blind reconstruction distinguishes reproducible interferon remodeling from less stable B-cell state assignments in systemic lupus erythematosus'

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest().upper()

def write(relative,text):
    path=PKG/relative; path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(text,encoding='utf-8',newline='\n')

def archive(files,path):
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in sorted(files.items()):
            info=zipfile.ZipInfo(name,(2026,9,9,0,0,0)); info.compress_type=zipfile.ZIP_DEFLATED
            z.writestr(info,data)

def manifest(files):
    buf=io.StringIO(newline=''); w=csv.writer(buf); w.writerow(['relative_path','bytes','sha256'])
    for name,data in sorted(files.items()): w.writerow([name,len(data),hashlib.sha256(data).hexdigest().upper()])
    return buf.getvalue().encode()

def prepare():
    PKG.mkdir(parents=True,exist_ok=True)
    mapping={}
    for ext in ('docx','pdf'):
        mapping['Manuscript/Manuscript.'+ext]=BASE/'documents'/('Manuscript_Scientific_Presentation_Freeze.'+ext)
        mapping['Supplementary_Information/Supplementary_Information.'+ext]=S7/'documents'/('Supplementary_Information_S7_readability.'+ext)
    for n in range(1,6):
        mapping[f'Figures/Figure_{n}.pdf']=next((BASE/'figures/figures').glob(f'Figure{n}_*.pdf'))
    for n in range(1,11):
        mapping[f'Author_Resources/Supplementary_Figures/Supplementary_Figure_S{n}.pdf']=next((BASE/'figures/figures').glob(f'Supplementary_Figure_S{n}_*.pdf'))
    for n,name in [(2,'Regulator_Sensitivity'),(3,'Full_Statistical_Results')]:
        mapping[f'Supplementary_Data/Supplementary_Data_{n}_{name}.zip']=OLD/'05_Source_Data'/(name+'.zip')
    provenance=[]
    for dest,source in mapping.items():
        target=PKG/dest; target.parent.mkdir(parents=True,exist_ok=True); shutil.copyfile(source,target)
        assert sha(target)==sha(source)
        provenance.append({'package_path':dest,'source_path':source.relative_to(ROOT).as_posix(),'sha256':sha(source)})
    figure_data={f.name:f.read_bytes() for f in (BASE/'figures/source_data').glob('*.csv')}
    assert len(figure_data)==15
    figure_data['MANIFEST_SHA256.csv']=manifest(figure_data)
    archive(figure_data,PKG/'Supplementary_Data/Supplementary_Data_1_Figure_Source_Data.zip')
    write('Author_Resources/Source_Provenance.json',json.dumps(provenance,indent=2))
    cover=f'''# Cover letter

9 September 2026

Editors, npj Systems Biology and Applications

Dear Editors,

Please consider our Article, "{TITLE}", for publication in npj Systems Biology and Applications.

The study addresses a systems-biology question in single-cell disease research: which biological layer remains reproducible when cell identity, composition and transcription are evaluated separately? Using public systemic lupus erythematosus datasets, we combined disease-blind B-lineage reconstruction, sample- or donor-level inference, assignment-uncertainty propagation and independent source-label-defined replication.

The contribution is an explicit hierarchy of reproducibility. A conventional-B-cell IFN/ISG program replicated despite weak genome-wide concordance, while end-to-end identity reconstruction and corrected external mapping retained clear limitations. STAT1/STAT2 analyses provided observational regulatory context, but broader interferon-response depletion weakened the discovery STAT2 result. We therefore make no claim of a universal cell-state taxonomy, causal regulator or validated clinical biomarker.

This separation of identity, abundance and molecular programs is relevant to the journal's readership in computational systems biology and systems immunology. The paper includes five main figures, supplementary information and three supporting data archives. Code and the version-specific scientific archive are available at https://github.com/1209433622cz-maker/sle-bcell-remodeling and https://doi.org/10.5281/zenodo.22151739; the supplied figure source data correspond to the current figures.

The work is original and is not under consideration elsewhere. Both authors have approved the scientific content and author contributions. The study received no funding, and the authors declare no competing interests. This secondary analysis used public, de-identified data and required no additional ethics approval. Generative AI assistance is disclosed in the manuscript.

Thank you for your consideration.

Sincerely,

Teng Qi, corresponding author, on behalf of both authors

School of Medicine, The Chinese University of Hong Kong, Shenzhen

Email: tengqi@link.cuhk.edu.cn

ORCID: https://orcid.org/0009-0007-7648-4776
'''
    author='''# Author information and declarations

## Author order

1. Zhi Chen, first author. Email: zhichen1@link.cuhk.edu.cn. ORCID: https://orcid.org/0009-0001-0072-5576. Academic title: MSc student in Bioinformatics.
2. Teng Qi, corresponding author. Email: tengqi@link.cuhk.edu.cn. ORCID: https://orcid.org/0009-0007-7648-4776. Academic title: MSc student in Bioinformatics.

Affiliation for both authors: School of Medicine, The Chinese University of Hong Kong, Shenzhen, Shenzhen 518172, China.

Correspondence: Teng Qi, MED Start-up Building, 2001 Longxiang Boulevard, Longgang District, Shenzhen 518172, China.

## First-author biography

Zhi Chen is an MSc student in Bioinformatics at The Chinese University of Hong Kong, Shenzhen. His research interests include integrated multi-omics analysis, clinical cancer research and the tumor microenvironment. He holds a BSc in Biomedical Sciences from Queen Mary University of London and an MB in Clinical Medicine from Nanchang University.

## Author contributions

Z.C.: Conceptualization, Data curation, Formal analysis, Investigation, Methodology, Software, Visualization, Writing - original draft. T.Q.: Conceptualization, Methodology, Project administration, Validation, Writing - review and editing.

## Declarations

Ethics: This secondary study used publicly available, de-identified human transcriptomic datasets. It involved no participant recruitment, intervention or collection of new specimens. No additional ethics approval was required. Ethics and consent procedures for the source studies are reported in the original publications.

Consent for publication: Not applicable; no identifiable participant information was used.

Funding: This study received no funding.

Competing interests: The authors declare no competing interests.

Acknowledgements: None.

Originality: The work is original and is not under consideration by another journal. Both authors confirmed the scientific content and author contributions in the project record.

Generative AI: OpenAI Codex was used for code drafting, workflow documentation, language editing and quality-control check development, as described in the manuscript. The authors remain responsible for the analyses, interpretations and final text. No AI system is an author.

Data and code: GEO GSE174188, GSE135779 and GSE23307; https://github.com/1209433622cz-maker/sle-bcell-remodeling; https://doi.org/10.5281/zenodo.22151739.
'''
    write('Cover_Letter/Cover_Letter.md',cover)
    write('Author_Resources/Authors_and_Declarations.md',author)
    for rel in ['Cover_Letter/Cover_Letter','Author_Resources/Authors_and_Declarations']:
        path=PKG/(rel+'.docx')
        docs.markdown_to_docx(PKG/(rel+'.md'),path,body_size=11,double_space=False,line_numbers=False,running_header='',title_override='Cover letter' if 'Cover_Letter' in rel else 'Author information and declarations')
        d=Document(path)
        for p in d.paragraphs:
            p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.space_before=Pt(0)
        d.core_properties.author='Zhi Chen; Teng Qi'; d.save(path)
    write('Author_Resources/Reporting_Information.md','''# Reporting information

Prepared information for the journal's current reporting-summary and editorial-policy forms; this file is not a completed official form.

Study: Secondary analysis of public, de-identified transcriptomic datasets. No new recruitment, random assignment or experimental intervention was performed. GSE174188 provides discovery/internal analyses, GSE135779 independent source-label-defined replication, and GSE23307 descriptive paired perturbational context (two donors).

Biological units: Sample-cohort strata for GSE174188 and donors for GSE135779; genes and cells are not independent donor replicates. Exact group sizes, exclusions, tests and multiplicity families are reported in Methods and Supplementary Tables S1, S3 and S7. Primary composition includes 43 control and 47 managed-SLE sample-cohort strata. No prospective sample-size calculation was performed for this retrospective analysis.

Blinding: Disease fields were withheld during identity adjudication; this does not imply that all later analyses were blinded. The corrected external mapping analysis was a post-outcome technical repair, not prospective validation.

Robustness: Frozen-representation and end-to-end resampling, assignment-exchange propagation, support thresholds, source-label omission and influence analyses are reported. The end-to-end B_ASC criterion and corrected external calibration failed; no corrected external disease effect was estimated. The paired two-donor perturbation is descriptive without inferential P values.

Statistics: Supplementary Table S7 specifies biological unit, test, sidedness, multiplicity and role. Full outputs and design matrices are in Supplementary Data 3. Regulator sensitivity outputs are in Supplementary Data 2. All current main and supplementary figure source tables are in Supplementary Data 1.

Data, software, source-study ethics, exclusions, covariate limitations, and AI assistance are specified in the manuscript Methods and availability statements. No new animal or human experiments were conducted for this study.
''')
    write('Supplementary_Data/Data_File_Descriptions.md','''# Supplementary data descriptions

Supplementary Data 1. Figure source data for Figures 1-5 and Supplementary Figures S1-S10, with a SHA-256 manifest. Tables correspond to the current panel assignments.

Supplementary Data 2. Regulator sensitivity outputs, including correlation-aware and IFN-overlap-depletion analyses, target-retention results and the original computational decision records.

Supplementary Data 3. Complete statistical results, sanitized design matrices, composition and transcription outputs, identity robustness analyses, and corrected external mapping calibration results, with provenance and integrity records.
''')
    write('README.md','''# SLE B-cell remodeling submission files

Target: npj Systems Biology and Applications. Article type: Article.

Upload Manuscript/Manuscript.docx (or its PDF), the five Figures/Figure_*.pdf files, Supplementary_Information/Supplementary_Information.pdf, the three Supplementary_Data ZIPs, and Cover_Letter/Cover_Letter.docx or PDF under the portal's corresponding categories. Select one format per document unless the portal requests both. The combined review PDF in Author_Resources is optional, not a second manuscript submission.

Author_Resources contains author metadata, the first-author biography, reporting information for portal forms, separately exported supplementary figures, and integrity/provenance files. These are preparation resources; do not upload the entire folder as supplementary research material. The ten supplementary figures are already embedded in the single Supplementary Information PDF.

Author order: Zhi Chen, then Teng Qi. Corresponding author: Teng Qi. Authors' scientific-content declarations are retained. The official reporting-summary/editorial-policy forms and account-level portal fields must be completed in the journal's current interface when requested; Reporting_Information.md is prepared supporting text, not a signed official form. This package has not been submitted.

The DOI 10.5281/zenodo.22151739 identifies the published scientific archive. Current presentation and figure-source assignments are supplied here and in GitHub; the DOI record has not been replaced by this packaging operation.

Official author information checked on 9 September 2026: https://www.nature.com/npjsba/for-authors-and-referees and https://www.nature.com/npjsba/for-authors-and-referees/editorial-process. These pages confirm online submission and flexible initial formatting. The detailed submission-guidelines page redirected to an inaccessible identity service during this check; no claim is made that current portal forms were completed.

Run `python Author_Resources/Verify_Package.py` from any directory to verify this extracted package. SHA256SUMS.csv lists all payload files; the ZIP checksum is supplied separately.
''')
    write('Author_Resources/Verify_Package.py', '''import csv, hashlib
from pathlib import Path
root=Path(__file__).resolve().parents[1]
with (root/'SHA256SUMS.csv').open(encoding='utf-8-sig') as f:
    rows=list(csv.DictReader(f))
assert {p.relative_to(root).as_posix() for p in root.rglob('*') if p.is_file()} == {r['relative_path'] for r in rows}|{'SHA256SUMS.csv'}
for r in rows:
    p=(root/r['relative_path']).resolve()
    assert p.is_relative_to(root.resolve())
    assert p.stat().st_size==int(r['bytes'])
    assert hashlib.sha256(p.read_bytes()).hexdigest().upper()==r['sha256']
print('PASS:',len(rows),'payload files verified')
''')

def finish():
    for row in json.loads((PKG/'Author_Resources/Source_Provenance.json').read_text()):
        assert sha(PKG/row['package_path'])==row['sha256']==sha(ROOT/row['source_path'])
    for p in PKG.glob('Supplementary_Data/*.zip'):
        with zipfile.ZipFile(p) as z:
            assert z.testzip() is None
            if 'MANIFEST_SHA256.csv' not in z.namelist():
                assert p.name=='Supplementary_Data_2_Regulator_Sensitivity.zip'
                assert sha(p)==sha(OLD/'05_Source_Data/Regulator_Sensitivity.zip')
                continue
            rows=list(csv.DictReader(io.StringIO(z.read('MANIFEST_SHA256.csv').decode('utf-8-sig'))))
            for row in rows:
                name=row.get('relative_path',row.get('path',row.get('file')))
                if name is None: raise ValueError(list(row))
                assert hashlib.sha256(z.read(name)).hexdigest().upper()==row['sha256'].upper()
    for p in [PKG/'Cover_Letter/Cover_Letter.pdf',PKG/'Author_Resources/Authors_and_Declarations.pdf']:
        assert p.exists()
    writer=PdfWriter()
    writer.append(PKG/'Manuscript/Manuscript.pdf')
    for n in range(1,6): writer.append(PKG/f'Figures/Figure_{n}.pdf')
    writer.write(PKG/'Author_Resources/Manuscript_With_Figures.pdf')
    assert len(PdfReader(PKG/'Author_Resources/Manuscript_With_Figures.pdf').pages)==36
    assert len(PdfReader(PKG/'Manuscript/Manuscript.pdf').pages)==31
    assert len(PdfReader(PKG/'Supplementary_Information/Supplementary_Information.pdf').pages)==15
    for p in PKG.rglob('*.docx'):
        assert 'Zhi Chen' in '\n'.join(x.text for x in Document(p).paragraphs) or 'Cover_Letter' in p.name
    files={p.relative_to(PKG).as_posix():p.read_bytes() for p in PKG.rglob('*') if p.is_file() and p.name!='SHA256SUMS.csv'}
    files['SHA256SUMS.csv']=manifest(files)
    (PKG/'SHA256SUMS.csv').write_bytes(files['SHA256SUMS.csv'])
    target=OUT/(PKG.name+'.zip')
    archive({PKG.name+'/'+n:b for n,b in files.items()},target)
    with zipfile.ZipFile(target) as z:
        assert z.testzip() is None
        assert all(z.read(PKG.name+'/'+n)==b for n,b in files.items())
    (OUT/'SHA256.txt').write_text(sha(target)+'  '+target.name+'\n')
    status={'status':'PACKAGE_BUILT_AND_HASH_VERIFIED','files':len(files),'zip_bytes':target.stat().st_size,'zip_sha256':sha(target),'main_pages':31,'supplement_pages':15,'main_figures':5,'supplementary_figures':10,'data_archives':3,'first_author':'Zhi Chen','corresponding_author':'Teng Qi','official_portal_forms_completed':False,'submitted':False}
    (OUT/'build_status.json').write_text(json.dumps(status,indent=2));print(json.dumps(status,indent=2))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--finish',action='store_true');a=p.parse_args()
    finish() if a.finish else prepare()
