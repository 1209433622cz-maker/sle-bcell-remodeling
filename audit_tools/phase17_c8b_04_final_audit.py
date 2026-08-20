#!/usr/bin/env python3
"""Audit and deterministically package the Gate C8B editorial preflight."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
RUN_DIR = ROOT / "phase17_v7" / "gateC8B" / "20260821_editorial_literature_preflight"
C8S_RUN = ROOT / "phase17_v7" / "gateC8S" / "20260821_supplementary_traceability_freeze"
PACKAGE = ROOT / "04_submission" / "package_genome_medicine_gateC8B_editorial_preflight_2026-08-21"
PACKAGE_ZIP = ROOT / "04_submission" / "package_genome_medicine_gateC8B_editorial_preflight_2026-08-21.zip"
MANUSCRIPT_MD = ROOT / "01_manuscript" / "manuscript_v13_genome_medicine_gateC8B_editorial_preflight_2026-08-21.md"
SUPPLEMENT_MD = ROOT / "01_manuscript" / "supplementary_information_v4_gateC8B_editorial_preflight_2026-08-21.md"
TARGET_MD = ROOT / "04_submission" / "journal_target_decision_gateC8B_2026-08-21.md"
COVER_MD = ROOT / "04_submission" / "cover_letter_genome_medicine_gateC8B_AUTHOR_COMPLETION_REQUIRED_2026-08-21.md"
REFERENCE_CSV = RUN_DIR / "references" / "reference_verification_gateC8B.csv"

MAIN_DOCX = PACKAGE / "main_text" / "Genome_Medicine_Manuscript_GateC8B_AUTHOR_COMPLETION_REQUIRED.docx"
SUPP_DOCX = PACKAGE / "additional_files" / "Additional_file_1_Supplementary_Information_GateC8B.docx"
COVER_DOCX = PACKAGE / "submission_docs" / "Genome_Medicine_Cover_Letter_GateC8B_AUTHOR_CONFIRMATION_REQUIRED.docx"
MAIN_PDF = PACKAGE / "internal_qc" / "wps_render_main" / "Genome_Medicine_Manuscript_GateC8B_AUTHOR_COMPLETION_REQUIRED_WPS.pdf"
SUPP_PDF = PACKAGE / "internal_qc" / "wps_render_supplement" / "Additional_file_1_Supplementary_Information_GateC8B_WPS.pdf"
COVER_PDF = PACKAGE / "internal_qc" / "wps_render_cover_letter" / "Genome_Medicine_Cover_Letter_GateC8B_AUTHOR_CONFIRMATION_REQUIRED_WPS.pdf"
TITLE = "Disease-blind single-cell reconstruction separates unstable B-cell states from reproducible interferon remodeling in systemic lupus erythematosus"
FIXED_ZIP_TIME = (2026, 8, 21, 0, 0, 0)
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def words(text: str) -> int:
    clean = re.sub(r"[`*_#|\[\]]", " ", text)
    return len(re.findall(r"\b[\w'-]+\b", clean))


def section(text: str, start: str, end: str | None) -> str:
    begin = text.index(start) + len(start)
    finish = text.index(end, begin) if end else len(text)
    return text[begin:finish].strip()


def ooxml(docx_path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(docx_path) as archive:
        return {name: archive.read(name) for name in archive.namelist() if name.endswith(".xml")}


def build_manifest() -> list[dict[str, object]]:
    manifest_path = PACKAGE / "MANIFEST_SHA256.csv"
    rows: list[dict[str, object]] = []
    for path in sorted(item for item in PACKAGE.rglob("*") if item.is_file() and item != manifest_path):
        rows.append(
            {
                "relative_path": path.relative_to(PACKAGE).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
        )
    with manifest_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["relative_path", "bytes", "sha256"])
        writer.writeheader()
        writer.writerows(rows)
    return rows


def write_deterministic_zip(output: Path) -> None:
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(item for item in PACKAGE.rglob("*") if item.is_file()):
            relative = (Path(PACKAGE.name) / path.relative_to(PACKAGE)).as_posix()
            info = zipfile.ZipInfo(relative, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            info.create_system = 3
            archive.writestr(info, path.read_bytes(), compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)


def build_deterministic_archive() -> tuple[int, str]:
    with tempfile.TemporaryDirectory(prefix="gateC8B_zip_") as temp_dir:
        first = Path(temp_dir) / "first.zip"
        second = Path(temp_dir) / "second.zip"
        write_deterministic_zip(first)
        write_deterministic_zip(second)
        if sha256(first) != sha256(second):
            raise RuntimeError("Deterministic package ZIP rebuild check failed")
        shutil.copyfile(first, PACKAGE_ZIP)
    return PACKAGE_ZIP.stat().st_size, sha256(PACKAGE_ZIP)


def package_readme(page_counts: dict[str, int], stats_hash: str) -> str:
    return f"""# Gate C8B Genome Medicine editorial and literature preflight

This package was generated on 21 August 2026 from the frozen Gate C8S scientific state.

## Status

- Editorial, literature, figure-wording, document and technical package: PASS.
- Scientific estimates changed: NO.
- Portal submission: BLOCKED pending author-controlled declarations, institutional ethics determination, repository licence and immutable archive DOI.
- Primary target: Genome Medicine; no journal-quartile assertion is frozen here.
- WPS render: {page_counts['main']} manuscript pages, {page_counts['supplement']} supplementary pages and {page_counts['cover']} cover-letter page.
- Main-panel assertions: PASS 46/46. Frozen supplementary-panel assertions: PASS 29/29.
- Figure 5c now labels E2F1, FOXM1, MYC and MYBL2 as proliferation specificity comparators.
- Sayadi et al. 2026 (doi:10.1016/j.jaut.2026.103575; PMID 42119160) is cited as external context, not independent replication.
- Full statistical archive remains the byte-identical Gate C8S freeze; SHA-256 `{stats_hash}`.

## Upload mapping after author completion

- `main_text/`: editable main manuscript.
- `figures/`: Figures 1-5 as PDF and 600-dpi PNG composites.
- `figures_supplementary/`: frozen Supplementary Figures S1-S7 as PDF and 600-dpi PNG composites.
- `additional_files/`: supplementary information, source data, regulator sensitivity and frozen full statistical results.
- `submission_docs/`: cover letter, author completion form, target decision and reporting checklist.
- `references/`: 27 DOI verification records and Vancouver-format metadata.
- `review_copies/`: WPS-rendered PDFs for author review.
- `internal_qc/`: accessibility reports, page rasters and audit evidence; do not upload unless requested.

## Hard stop

Do not submit until all visible author-controlled items have been completed and verified. Do not modify frozen scientific estimates during Gate C8B.
"""


def main() -> None:
    manuscript = MANUSCRIPT_MD.read_text(encoding="utf-8")
    supplement = SUPPLEMENT_MD.read_text(encoding="utf-8")
    cover = COVER_MD.read_text(encoding="utf-8")
    target = TARGET_MD.read_text(encoding="utf-8")
    main_status = read_json(RUN_DIR / "01_FIGURE_BUILD_STATUS.json")
    main_assertions = read_json(RUN_DIR / "02_PANEL_DATA_ASSERTIONS.json")
    reference_status = read_json(RUN_DIR / "01_GATE_C8B_REFERENCE_STATUS.json")
    source_status = read_json(RUN_DIR / "03_GATE_C8B_SOURCE_BUILD_STATUS.json")
    document_status = read_json(RUN_DIR / "04_GATE_C8B_DOCUMENT_BUILD_STATUS.json")
    supp_status = read_json(C8S_RUN / "04_SUPPLEMENTARY_FIGURE_BUILD_STATUS.json")
    supp_assertions = read_json(C8S_RUN / "03_SUPPLEMENTARY_PANEL_DATA_ASSERTIONS.json")
    stats_status = read_json(C8S_RUN / "05_FULL_STATISTICAL_RESULTS_STATUS.json")

    checks: dict[str, dict[str, object]] = {}

    def check(name: str, passed: bool, detail: object) -> None:
        checks[name] = {"pass": bool(passed), "detail": str(detail)}

    main_rows = main_assertions.get("checks", [])
    check(
        "main_figure_assertions",
        main_status.get("status") == "C8B_MAIN_FIGURES_BUILT_WITH_ASSERTIONS"
        and main_assertions.get("status") == "PASS"
        and len(main_rows) == 46
        and all(row.get("pass") is True for row in main_rows),
        f"passed={sum(row.get('pass') is True for row in main_rows)}/46",
    )
    supp_rows = supp_assertions.get("checks", [])
    check(
        "supplementary_figure_assertions",
        supp_status.get("status") == "PASS_GATE_C8S_SUPPLEMENTARY_FIGURES_BUILT"
        and supp_assertions.get("status") == "PASS"
        and len(supp_rows) == 29
        and all(row.get("pass") is True for row in supp_rows),
        f"frozen C8S assertions passed={sum(row.get('pass') is True for row in supp_rows)}/29",
    )

    figure_details = []
    figures_pass = True
    for prefix, directory, total in (
        ("Figure", PACKAGE / "figures", 5),
        ("Supplementary_Figure_S", PACKAGE / "figures_supplementary", 7),
    ):
        for number in range(1, total + 1):
            pngs = list(directory.glob(f"{prefix}{number}_*.png"))
            pdfs = list(directory.glob(f"{prefix}{number}_*.pdf"))
            ok = len(pngs) == 1 and len(pdfs) == 1
            dimensions = (0, 0)
            if ok:
                with Image.open(pngs[0]) as image:
                    dimensions = image.size
                ok = dimensions[0] >= 4000 and dimensions[1] >= 3000 and pngs[0].stat().st_size > 0 and pdfs[0].stat().st_size > 0
            figures_pass &= ok
            figure_details.append({"figure": f"{prefix}{number}", "pixels": dimensions, "pass": ok})
    check("publication_figure_files", figures_pass, f"pairs={len(figure_details)}; all >=4000x3000 and non-empty")

    figure5_pdf = next((PACKAGE / "figures").glob("Figure5_*.pdf"))
    figure5_text = " ".join((page.extract_text() or "") for page in PdfReader(figure5_pdf).pages)
    figure5_text = re.sub(r"\s+", " ", figure5_text)
    comparator_phrase = "Prespecified proliferation specificity comparators"
    check(
        "figure5c_specificity_wording",
        comparator_phrase.lower() in figure5_text.lower()
        and "Prespecified proliferation controls".lower() not in figure5_text.lower(),
        comparator_phrase,
    )

    source_dir = PACKAGE / "additional_files" / "source_data"
    source_rows = read_csv(source_dir / "SHA256SUMS.csv")
    source_csvs = sorted(source_dir.glob("*Figure*_source_data.csv"))
    source_hashes_ok = len(source_csvs) == 12 and len(source_rows) == 12 and all(
        sha256(source_dir / row["file"]) == row["sha256"].upper() for row in source_rows
    )
    fig5_rows = read_csv(source_dir / "Figure5_source_data.csv")
    panel_counts = {panel: sum(row["panel"] == panel for row in fig5_rows) for panel in ("B", "C", "D", "E")}
    mapping_ok = (
        panel_counts == {"B": 12, "C": 12, "D": 3, "E": 2}
        and all(row["series"] == "MSigDB_M5911_NES" for row in fig5_rows if row["panel"] == "D")
        and all(row["series"] == "GSE23307_mean_paired_log2p1_effect" for row in fig5_rows if row["panel"] == "E")
    )
    check("figure_source_data", source_hashes_ok and mapping_ok, f"csv=12; Figure5 panels={panel_counts}")

    stats_zip = PACKAGE / "additional_files" / "Additional_file_4_Full_Statistical_Results_GateC8S_FROZEN.zip"
    with zipfile.ZipFile(stats_zip) as archive:
        stats_entries = [name for name in archive.namelist() if not name.endswith("/")]
    stats_hash = sha256(stats_zip)
    stats_ok = (
        stats_status.get("status") == "PASS_GATE_C8S_FULL_STATISTICAL_ARCHIVE_BUILT"
        and len(stats_entries) == 63
        and stats_hash == stats_status.get("archive_sha256")
        and stats_status.get("complete_gene_result_files") == 12
        and stats_status.get("sanitized_design_matrices") == 12
    )
    check("frozen_statistical_archive", stats_ok, f"entries={len(stats_entries)}; sha256={stats_hash}")

    abstract_words = words(section(manuscript, "## Abstract", "## Keywords"))
    refs = [int(value) for value in re.findall(r"^(\d+)\.\s", section(manuscript, "## References", None), flags=re.M)]
    required_tokens = [
        f"# {TITLE}",
        "**Version:** Gate C8B editorial and literature preflight v13",
        "contextual rather than replication evidence",
        "external biological context rather than independent validation",
        "doi:10.1016/j.jaut.2026.103575",
        "Prespecified E2F1, FOXM1, MYC and MYBL2 proliferation specificity comparators",
        "Gate C8S remains the canonical frozen scientific state",
    ]
    missing_tokens = [token for token in required_tokens if token not in manuscript]
    check(
        "manuscript_editorial_literature_contract",
        not missing_tokens
        and abstract_words == 314
        and refs == list(range(1, 32))
        and "### Authors' information" not in manuscript,
        f"abstract={abstract_words}; references={len(refs)}; missing={missing_tokens or 'none'}; Authors information absent",
    )
    check(
        "cover_letter_scope",
        "an interferon-responsive transcriptional program" in cover and "a type I IFN/ISG program" not in cover,
        "ligand-agnostic interferon-responsive wording present",
    )
    check(
        "reference_verification",
        reference_status.get("decision") == "PASS"
        and reference_status.get("doi_records") == 27
        and len(read_csv(REFERENCE_CSV)) == 27
        and all(row["status"] == "PASS" for row in read_csv(REFERENCE_CSV)),
        "Crossref PASS=27/27; manuscript references=31; PMID 42119160 recorded",
    )
    check(
        "source_build_contract",
        source_status.get("status") == "PASS_GATE_C8B_EDITORIAL_LITERATURE_SOURCES_BUILT"
        and source_status.get("scientific_estimates_changed") is False
        and supplement.count("[[SUPPLEMENTARY_FIGURE:S") == 7,
        "v13 source PASS; 7 supplement figure markers; estimates unchanged",
    )
    check("journal_target", "**Primary submission: Genome Medicine.**" in target, "Genome Medicine retained without a frozen quartile assertion")

    docx_sizes = {path.name: path.stat().st_size if path.exists() else 0 for path in (MAIN_DOCX, SUPP_DOCX, COVER_DOCX)}
    assets = document_status.get("assets", {})
    check(
        "document_build_status",
        document_status.get("status") == "PASS_GATE_C8B_DOCUMENTS_BUILT"
        and len(document_status.get("outputs", [])) == 3
        and assets.get("main_figure_files") == 10
        and assets.get("supplementary_figure_files") == 14
        and assets.get("source_files") == 12
        and all(size > 30_000 for size in docx_sizes.values()),
        docx_sizes,
    )

    main_xml = ooxml(MAIN_DOCX)
    main_document = ET.fromstring(main_xml["word/document.xml"])
    main_settings = ET.fromstring(main_xml["word/settings.xml"])
    main_styles = ET.fromstring(main_xml["word/styles.xml"])
    line_numbers = main_document.find(".//w:sectPr/w:lnNumType", NS) is not None
    page_field = b"PAGE" in b"".join(value for key, value in main_xml.items() if key.startswith("word/footer"))
    even_odd = main_settings.find(".//w:evenAndOddHeaders", NS) is not None
    normal_style = main_styles.find(".//w:style[@w:styleId='Normal']", NS)
    spacing = normal_style.find(".//w:spacing", NS) if normal_style is not None else None
    double_spacing = spacing is not None and spacing.get(f"{{{NS['w']}}}line") == "480"
    check("main_docx_ooxml", line_numbers and page_field and even_odd and double_spacing, f"line_numbers={line_numbers}; page_field={page_field}; even_odd={even_odd}; double_spacing={double_spacing}")

    supplement_doc = Document(SUPP_DOCX)
    geometry = []
    for table in supplement_doc.tables:
        properties = table._tbl.find(qn("w:tblPr"))
        grid = table._tbl.find(qn("w:tblGrid"))
        widths = []
        for row in table.rows:
            for cell in row.cells:
                cell_properties = cell._tc.find(qn("w:tcPr"))
                widths.append(cell_properties is not None and cell_properties.find(qn("w:tcW")) is not None)
        geometry.append(
            properties is not None
            and properties.find(qn("w:tblW")) is not None
            and properties.find(qn("w:tblInd")) is not None
            and grid is not None
            and all(widths)
        )
    check("supplement_docx_structure", len(geometry) == 8 and all(geometry) and len(supplement_doc.inline_shapes) == 7, f"tables={len(geometry)}; explicit_geometry={sum(geometry)}; inline_figures={len(supplement_doc.inline_shapes)}")

    a11y_details = []
    for name in ("main_text_a11y.json", "supplement_a11y.json", "cover_letter_a11y.json"):
        report = read_json(PACKAGE / "internal_qc" / name)
        counts = report.get("counts", {})
        a11y_details.append({"file": name, **counts})
    check("docx_accessibility", all(sum(int(row.get(level, 0)) for level in ("high", "medium", "low")) == 0 for row in a11y_details), a11y_details)

    render_details = []
    page_counts: dict[str, int] = {}
    renders_ok = True
    for name, pdf, minimum, maximum in (
        ("main", MAIN_PDF, 25, 30),
        ("supplement", SUPP_PDF, 12, 13),
        ("cover", COVER_PDF, 1, 1),
    ):
        reader = PdfReader(pdf) if pdf.exists() else None
        pages = len(reader.pages) if reader else 0
        pngs = len(list(pdf.parent.glob("final-page-*.png")))
        ok = pdf.exists() and pdf.stat().st_size > 30_000 and minimum <= pages <= maximum and pngs == pages
        renders_ok &= ok
        page_counts[name] = pages
        render_details.append({"document": name, "pages": pages, "page_pngs": pngs, "bytes": pdf.stat().st_size if pdf.exists() else 0})
    check("wps_render_integrity", renders_ok, render_details)

    author_tokens = [
        "Zhi Chen [1] and Teng Qi [1,*]",
        "zhichen1@link.cuhk.edu.cn",
        "tengqi@link.cuhk.edu.cn",
        "0009-0001-0072-5576",
        "0009-0007-7648-4776",
        "School of Medicine, The Chinese University of Hong Kong, Shenzhen",
    ]
    check("author_identity", all(token in manuscript for token in author_tokens), "names, emails, ORCIDs and affiliation present")

    hard_stops = [
        "institutional ethics determination",
        "competing interests",
        "funding",
        "final CRediT contributions and all-author approval",
        "acknowledgements",
        "all-author originality/submission confirmation",
        "open-source licence and immutable archive DOI",
    ]
    manuscript_placeholders = manuscript.count("[[")
    cover_placeholders = cover.count("[[")
    check("hard_stops_visible", manuscript_placeholders == 6 and cover_placeholders == 2, f"manuscript placeholders={manuscript_placeholders}; cover placeholders={cover_placeholders}")

    technical_checks = [name for name in checks if name != "hard_stops_visible"]
    technical_pass = all(checks[name]["pass"] for name in technical_checks)
    decision = (
        "PASS_GATE_C8B_EDITORIAL_LITERATURE_PREFLIGHT_AUTHOR_ACTION_REQUIRED"
        if technical_pass
        else "HOLD_GATE_C8B_TECHNICAL_REPAIR_REQUIRED"
    )
    audit = {
        "created_at": "2026-08-21",
        "decision": decision,
        "primary_target": "Genome Medicine",
        "source_scientific_freeze": "Gate C8S",
        "scientific_estimates_changed": False,
        "technical_package_pass": technical_pass,
        "portal_submission_authorized": False,
        "checks": checks,
        "hard_stops": hard_stops,
        "next_stage": "Complete Gate C8B author-controlled declarations, institutional ethics determination, repository licence and immutable archive DOI; then replace placeholders, run one final WPS review and portal preflight without reopening scientific analysis",
    }

    audit_json = RUN_DIR / "05_GATE_C8B_FINAL_AUDIT.json"
    audit_md = RUN_DIR / "05_GATE_C8B_FINAL_AUDIT.md"
    audit_json.write_text(json.dumps(audit, indent=2) + "\n", encoding="utf-8", newline="\n")
    lines = [
        "# Gate C8B editorial and literature preflight audit",
        "",
        f"**Decision:** `{decision}`",
        "",
        f"**Technical package:** {'PASS' if technical_pass else 'HOLD'}",
        "",
        "**Scientific estimates changed:** NO",
        "",
        "**Portal submission authorized:** NO - author-controlled hard stops remain.",
        "",
        "## Checks",
        "",
        "| Check | Result | Detail |",
        "|---|---:|---|",
    ]
    for name, result in checks.items():
        detail = str(result["detail"]).replace("|", "/").replace("\n", " ")
        lines.append(f"| `{name}` | {'PASS' if result['pass'] else 'FAIL'} | {detail} |")
    lines.extend(
        ["", "## Author-controlled hard stops", ""]
        + [f"- {item}" for item in hard_stops]
        + ["", "## Next stage", "", audit["next_stage"]]
    )
    audit_md.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")

    review_dir = PACKAGE / "review_copies"
    review_dir.mkdir(parents=True, exist_ok=True)
    for path in (MAIN_PDF, SUPP_PDF, COVER_PDF):
        shutil.copy2(path, review_dir / path.name)
    (PACKAGE / "README_GATE_C8B_PACKAGE.md").write_text(
        package_readme(page_counts, stats_hash), encoding="utf-8", newline="\n"
    )
    shutil.copy2(audit_json, PACKAGE / "internal_qc" / "GATE_C8B_FINAL_AUDIT.json")
    shutil.copy2(audit_md, PACKAGE / "internal_qc" / "GATE_C8B_FINAL_AUDIT.md")
    manifest_rows = build_manifest()
    shutil.copy2(PACKAGE / "MANIFEST_SHA256.csv", RUN_DIR / "06_GATE_C8B_INTEGRITY_MANIFEST.csv")
    zip_bytes, zip_hash = build_deterministic_archive()
    summary = {
        **audit,
        "deterministic_archive_rebuild_match": True,
        "package_files_manifested": len(manifest_rows),
        "page_counts": page_counts,
        "package_zip": PACKAGE_ZIP.relative_to(ROOT).as_posix(),
        "package_zip_bytes": zip_bytes,
        "package_zip_sha256": zip_hash,
    }
    (RUN_DIR / "07_GATE_C8B_PACKAGE_STATUS.json").write_text(
        json.dumps(summary, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    print(json.dumps(summary, indent=2))
    if not technical_pass:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
