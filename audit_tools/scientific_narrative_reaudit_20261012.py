"""Source-bound scientific narrative candidate and read-only evidence audit."""
from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Pt

DATE = "2026-10-12"
SKILL_NAME = "QiTeng_Academic_Writing_Skill_v0.3.24.2_CompressionGovernance_Release_2026-09-20"


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def save_json(path, value):
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def save_csv(path, rows, fields=None):
    fields = fields or list(rows[0])
    with Path(path).open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def paragraph(text, prefix):
    hits = [p for p in text.split("\n\n") if p.startswith(prefix)]
    if len(hits) != 1:
        raise ValueError(f"Expected unique paragraph: {prefix}: {len(hits)}")
    return hits[0]


def patch_docx(source, target, replacements, insertions):
    doc = Document(source)
    for old, new in replacements.items():
        hits = [p for p in doc.paragraphs if p.text == old]
        if len(hits) != 1:
            raise ValueError(f"DOCX paragraph not unique: {old[:70]} ({len(hits)})")
        p = hits[0]
        props = copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
        p.clear()
        r = p.add_run(new)
        if props is not None:
            r._r.insert(0, props)
    for anchor, new, style in insertions:
        hits = [p for p in doc.paragraphs if p.text == anchor]
        if len(hits) != 1:
            raise ValueError(f"DOCX insertion anchor not unique: {anchor[:70]}")
        template = hits[0] if style is None else next(q for q in doc.paragraphs if q.style.name == style and q.text)
        p = hits[0].insert_paragraph_before(new, style=template.style)
        if template.runs and template.runs[0]._r.rPr is not None:
            p.runs[0]._r.insert(0, copy.deepcopy(template.runs[0]._r.rPr))
    doc.save(target)
    with zipfile.ZipFile(source) as a, zipfile.ZipFile(target) as b:
        am = {n: hashlib.sha256(a.read(n)).hexdigest() for n in a.namelist() if n.startswith("word/media/")}
        bm = {n: hashlib.sha256(b.read(n)).hexdigest() for n in b.namelist() if n.startswith("word/media/")}
        if am != bm:
            raise ValueError("DOCX media changed")
    return len(am)


def docx_markdown(source, root):
    doc = Document(source)
    blocks, reference_number, figure_number = [], 0, None
    for element in doc.element.body:
        if element.tag.endswith("}p"):
            p = Paragraph(element, doc)
            text = p.text.strip()
            if text:
                match = re.search(r"Supplementary Figure S(\d+) \|", text)
                if match:
                    figure_number = int(match.group(1))
                if p.style.name == "Title":
                    text = "# " + text
                elif p.style.name.startswith("Heading "):
                    text = "#" * (int(p.style.name.split()[-1]) + 1) + " " + text
                elif p.style.name == "Reference":
                    reference_number += 1
                    text = str(reference_number) + ". " + text
                blocks.append(text)
            if p._p.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"):
                if figure_number is None:
                    raise ValueError("Image without supplementary figure owner")
                matches = list((root / "4_figures/Supplementary").glob(f"Supplementary_Figure_S{figure_number}_*.png"))
                if len(matches) != 1:
                    raise ValueError("Supplementary image not unique")
                blocks.append(f"![Supplementary Figure S{figure_number}](../../{matches[0].relative_to(root).as_posix()})")
        elif element.tag.endswith("}tbl"):
            table = Table(element, doc)
            rows = [[c.text.replace("\n", " ").replace("|", "\\|") for c in r.cells] for r in table.rows]
            blocks.append("\n".join(["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])] + ["| " + " | ".join(row) + " |" for row in rows[1:]]))
    return "\n\n".join(blocks) + "\n", reference_number


def build(root, out, report):
    main = (root / "01_manuscript/Manuscript.md").read_text(encoding="utf-8")
    supp = (root / "01_manuscript/Supplementary_Information.md").read_text(encoding="utf-8")
    replacements = {}
    ledger = []

    def edit(prefix, new, reason):
        old = paragraph(main, prefix)
        replacements[old] = new
        ledger.append(dict(object="Manuscript", operation="replace", reason=reason, before=old, after=new))
        return old

    abstract = ("Single-cell disease studies can conflate identity, abundance and transcription when unstable annotations are treated as fixed states. We reanalysed public systemic lupus erythematosus datasets using disease-blind B-lineage reconstruction and biological-unit-aware inference. Among 150,402 discovery B-lineage cells, end-to-end resampling failed the antibody-secreting-cell overlap criterion, restricting broad compartments to an analysis scaffold. The primary abundance contrast lacked support (odds ratio 0.947, 95% confidence interval 0.636-1.410). In 89 conventional-B pseudobulks, the frozen IFN/ISG program was higher in source-defined managed SLE (adjusted score difference 0.837, 95% confidence interval 0.525-1.148; four-program q=2.98 x 10^-6). It replicated in 43 independent GSE135779 childhood donors using a source-label-defined broad-B analogue (difference 1.042, 95% confidence interval 0.681-1.402), despite weak genome-wide concordance (rho=0.026). Corrected source-label-independent remapping failed calibration; no corrected disease effect was estimated. STAT1/STAT2 target analyses converged observationally but weakened after broader interferon-gene depletion. Thus, the prespecified process-level interferon association retained support where tested hard-state assignments did not, without establishing a universal taxonomy, causal regulator or clinical utility.")
    edit("Single-cell disease studies can", abstract, "Restore quantitative primary/external anchors and denominators; no added biological claim.")
    old_program = paragraph(main, "Other prespecified programs were less consistent.")
    edit("Other prespecified programs were less consistent.", old_program.replace("(Fig. 3d).", "(Fig. 3a)."), "Four-program competition belongs to 3a, not specificity panel 3d.")
    specificity = ("Prespecified specificity controls did not show the same cross-contrast pattern. Platelet/ambient and ASC/UPR estimates had intervals spanning zero in both the primary and donor-nonoverlap analyses. The pan-B control was positive in the primary analysis (effect 0.409, 95% confidence interval 0.119-0.699; nominal P=0.00624) but not in donor-nonoverlap replication (effect -0.194, 95% confidence interval -0.835 to 0.447). These diagnostics distinguish the replicated IFN program from a uniformly recurring technical or pan-B pattern; they do not prove the absence of contamination or all nonspecific expression effects (Fig. 3d).")
    edit("In a later, outcome-exposed sensitivity analysis", "The later measurement-rule analysis addresses a separate uncertainty: how an observed association depends on the expression-mixing operator. Its two intervals describe the same 68-donor contrast, not two replications. Because its fixed-reference T0 score and population differ from the primary TMM-score analysis, their effect magnitudes cannot be combined or interpreted as a proportion explained by composition. Synthetic G2C qualification supports the tested interval procedure only within its generating regimes; it does not establish that all capture, misclassification or mixing mechanisms in the real data were modelled correctly (Supplementary Table S10).", "Remove new numerical evidence from Discussion; retain interpretive delta and scale/chronology boundaries.")
    measurement = ("A later, explicitly outcome-exposed measurement-rule sensitivity concerned an already observed contrast rather than an additional replication. In the same 68 donors (31 controls and 37 source-defined managed-SLE donors), fixed-reference profile mixing minus pooled-count scoring yielded a disease-coefficient difference of 0.034925 T0 units. The earlier paired-percentile bootstrap interval was 0.013789-0.057130; the subsequent G2C-qualified HC3-t(64) interval was 0.011834-0.058016 (nominal P=0.003612). Both procedures use the same outcomes and estimand. This difference is neither a mediated fraction nor a cell-intrinsic disease effect, and its scale is distinct from the primary TMM-score model (Supplementary Table S10).")
    # Keep the later measurement question separate from observational regulator evidence.
    program_new = replacements[old_program]
    insertions = [(program_new, specificity, None)]
    measurement_heading = "Later measurement-rule sensitivity re-estimates uncertainty for the same contrast"
    insertions.append(("Discussion", measurement_heading, "Heading 2"))
    insertions.append(("Discussion", measurement, "Normal"))
    for text, reason in [(specificity, "Provide 3d with an explicit Results owner, including the positive primary pan-B control."), (measurement, "Methods-Results mirror: report the existing outcome-exposed estimate before discussing it.")]:
        ledger.append(dict(object="Manuscript", operation="insert", reason=reason, before="", after=text))
    ledger.append(dict(object="Manuscript Results", operation="section_boundary", reason="Place later measurement sensitivity after the regulatory closing paragraph under its own bounded heading; do not imply it is an additional regulatory evidence layer.", before="no separate measurement Results owner", after=measurement_heading))
    source_doc = Document(root / "5_manuscript/current/Manuscript.docx")
    aliases = {}
    for prefix in ["For each tested Ensembl feature,", "a, Standardized IFN/ISG effects for childhood", "a, Evidence classes and"]:
        hits = [p.text for p in source_doc.paragraphs if p.text.startswith(prefix)]
        if len(hits) != 1:
            raise ValueError("Expected unique approved DOCX paragraph")
        md_prefix = "a, Evidence classes and interpretive roles." if prefix == "a, Evidence classes and" else prefix
        old_md = paragraph(main, md_prefix)
        aliases[old_md] = hits[0]
        ledger.append(dict(object="MD/DOCX baseline reconciliation", operation="synchronize", reason="Use current DOCX's defined abbreviations and reader-facing wording; retain a separate record of pre-existing drift.", before=old_md, after=hits[0]))
    old = aliases[paragraph(main, "For each tested Ensembl feature,")]
    edit("For each tested Ensembl feature,", old + " These target-model intervals quantify the fitted gene-statistic slope, not donor-level TF activity or new biological replication; correlation-aware sensitivities address the dependence among targets.", "Declare target-model uncertainty unit; preserve original 24-test family.")
    old = aliases[paragraph(main, "a, Standardized IFN/ISG effects for childhood")]
    edit("a, Standardized IFN/ISG effects for childhood", old + " Scores were standardized within each analysis; the displayed estimates are not pooled on a common absolute-expression scale. Combined and childhood analyses overlap and are not separate independent replications.", "Prevent visual common-axis display being mistaken for meta-analysis or two independent cohorts.")
    old = aliases[paragraph(main, "a, Evidence classes and interpretive roles.")]
    edit("a, Evidence classes and interpretive roles.", old + " Intervals in b-d describe signed-target ULM slopes fitted to gene-ranked statistics, not participant-level TF activity. The original 24-test family and branch-specific six-test sensitivity families remain distinct.", "Make gene/participant units and multiplicity visible in the regulatory legend.")
    old = paragraph(main, "The datasets analysed are publicly available")
    edit("The datasets analysed are publicly available", "The datasets analysed are publicly available through NCBI GEO under GSE174188, GSE135779 and GSE23307 [17,18,20]. Figure source data and complete statistical results for the core analysis are supplied as Supplementary Data 1-3 and archived at https://doi.org/10.5281/zenodo.22151739 [33]. The later measurement-rule sensitivity is reported separately in Supplementary Table S10 and its accompanying source-data file; it is not part of that earlier version-specific archive. Large recomputable matrices are not duplicated from their source repositories. Third-party GEO and CELLxGENE data remain subject to their source terms.", "Separate the existing core DOI from later RP v4 sensitivity evidence; do not imply that the earlier archive contains new results.")
    edit("Analysis code, executable decision records", "Core analysis code, decision records, environments and restoration instructions are available at https://github.com/1209433622cz-maker/sle-bcell-remodeling (release v1.1.0; scientific content commit f1859ff8498d5569a1d5027b36ed18c8b7c7536f) and archived at https://doi.org/10.5281/zenodo.22151739 [33]. The same repository retains the later measurement-rule sensitivity code and write-once result record under the RP v4 branch of the analysis; the Stage I result records its execution binding to commit b413ce09fae520eef0a7a781ee71ddece7606be0. Original project code is licensed under the MIT License.", "Provide the execution-code binding for the later sensitivity without changing a DOI or release.")
    candidate = main
    for old, new in replacements.items():
        if candidate.count(old) != 1:
            raise ValueError("Markdown replacement not unique")
        candidate = candidate.replace(old, new)
    for anchor, new, style in insertions:
        md_anchor = "## Discussion" if anchor == "Discussion" else anchor
        md_new = "### " + new if style == "Heading 2" else new
        candidate = candidate.replace(md_anchor, md_new + "\n\n" + md_anchor, 1)
    (out / "Manuscript.md").write_text(candidate, encoding="utf-8")
    old = paragraph(supp, "This single Supplementary Information file")
    new = old.replace("Tables S1-S9", "Tables S1-S10")
    (out / "Supplementary_Information.md").write_text(supp.replace(old, new), encoding="utf-8")
    ledger.append(dict(object="Supplementary Information", operation="replace", reason="Include the actually integrated Table S10 in the overview.", before=old, after=new))
    source = root / "5_manuscript/current"
    docx_replacements = {aliases.get(old, old): new for old, new in replacements.items()}
    media_main = patch_docx(source / "Manuscript.docx", out / "Manuscript.docx", docx_replacements, insertions)
    media_supp = patch_docx(source / "Supplementary_Information.docx", out / "Supplementary_Information.docx", {old: new}, [])
    editable = Document(out / "Manuscript.docx")
    for prefix in ["The later measurement-rule analysis", "A later, explicitly outcome-exposed measurement-rule sensitivity used"]:
        hits = [p for p in editable.paragraphs if p.text.startswith(prefix)]
        if len(hits) != 1:
            raise ValueError("Measurement-rule body paragraph not unique")
        for run in hits[0].runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
    editable.save(out / "Manuscript.docx")
    ledger.append(dict(object="Manuscript Methods/Discussion", operation="source_style_rebuild", reason="Correct two inherited 9 pt Stage K body paragraphs to the manuscript's 12 pt body standard; no content change.", before="9 pt body additions", after="Times New Roman 12 pt"))
    supplementary = Document(out / "Supplementary_Information.docx")
    s1 = [i for i, p in enumerate(supplementary.paragraphs) if p.text.startswith("Supplementary Figure S1 |")]
    if len(s1) != 1:
        raise ValueError("S1 heading not unique")
    start = s1[0]
    for p in supplementary.paragraphs[start:]:
        if p._p.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"):
            break
        p.paragraph_format.keep_with_next = True
    supplementary.save(out / "Supplementary_Information.docx")
    ledger.append(dict(object="Supplementary Figure S1", operation="source_pagination_rebuild", reason="Bind S1 heading/legend to its unchanged image after Table S10 integration; correct inherited page 6/7 separation.", before="heading/legend on page 6; image on page 7", after="keep-with-next through image anchor"))
    synced, reference_count = docx_markdown(out / "Manuscript.docx", root)
    synced_supp, _ = docx_markdown(out / "Supplementary_Information.docx", root)
    baseline_export, _ = docx_markdown(source / "Manuscript.docx", root)
    (report / "BASELINE_DOCX_TEXT.md").write_text(baseline_export, encoding="utf-8")
    import difflib
    (report / "PREEXISTING_MD_DOCX_DRIFT.diff").write_text("\n".join(difflib.unified_diff(main.splitlines(), baseline_export.splitlines(), fromfile="canonical_Manuscript.md", tofile="current_DOCX_text")) + "\n", encoding="utf-8")
    (out / "Manuscript.md").write_text(synced, encoding="utf-8")
    (out / "Supplementary_Information.md").write_text(synced_supp, encoding="utf-8")
    ledger.append(dict(object="All candidate Markdown", operation="source_sync", reason="Export the final editable candidate including its current title and abbreviation expansions; reference numbering resolves OOXML automatic lists. No scientific number selection.", before=main.splitlines()[0], after=synced.splitlines()[0]))
    save_csv(report / "QITENG_EDIT_LEDGER.csv", ledger)
    save_json(report / "CANDIDATE_BUILD.json", dict(abstract_words=len(abstract.split()), changes=len(ledger), references=reference_count, main_media=media_main, supp_media=media_supp, canonical_overwritten=False, editable_text_authority="5_manuscript/current/*.docx", markdown_generated_from_candidate_docx=True))
    if main.split("## References", 1)[1].split("## Figure legends", 1)[0] != candidate.split("## References", 1)[1].split("## Figure legends", 1)[0]:
        raise ValueError("Reference bibliography changed")
    return candidate


def audit(root, report):
    checks = []
    numeric_code = """import csv, json, sys
from scipy import stats
def rows(p):
    with open(p, encoding='utf-8-sig', newline='') as f: return list(csv.DictReader(f))
a = [r for r in rows(sys.argv[1]) if r['panel'] == 'a']
g = [r for r in rows(sys.argv[2]) if r['panel'] == 'c']
s = json.load(open(sys.argv[3], encoding='utf-8'))
result = {'ci': {r['program_id']: [float(r['effect'])-stats.t.ppf(.975,float(r['df']))*float(r['se_hc3']),float(r['effect'])+stats.t.ppf(.975,float(r['df']))*float(r['se_hc3'])] for r in a}, 'rho': stats.spearmanr([float(r['gse174188_logFC']) for r in g],[float(r['gse135779_logFC']) for r in g]).statistic, 'p': 2*stats.t.sf(abs(s['test_statistic']),s['df'])}
print(json.dumps(result))
"""
    number_run = subprocess.run(["D:/bioinfor/python.exe", "-c", numeric_code, str(root / "4_figures/Fig3/source_data/Figure_3_source_data.csv"), str(root / "4_figures/Fig4/source_data/Figure_4_source_data.csv"), str(root / "3_results/current/StageI_Result.json")], capture_output=True, text=True, check=True)
    numbers = json.loads(number_run.stdout)
    save_json(report / "INDEPENDENT_NUMERIC_RECOMPUTATION.json", numbers)
    def check(name, ok, detail):
        checks.append(dict(check=name, status="PASS" if ok else "HOLD", detail=str(detail)))
    assets = []
    figures = []
    for number in range(1, 6):
        figures.append(root / f"4_figures/Fig{number}/canonical/Figure_{number}.pdf")
    figures += sorted((root / "4_figures/Supplementary").glob("*.pdf"))
    for pdf in figures:
        with fitz.open(pdf) as doc:
            sizes = [s["size"] for p in doc for b in p.get_text("dict")["blocks"] if "lines" in b for l in b["lines"] for s in l["spans"] if s["text"].strip()]
            fonts = [font for p in doc for font in p.get_fonts(full=True)]
            embedded = all(f[2] != "Type3" and doc.extract_font(f[0])[3] for f in fonts if f[0] > 0)
            widths = [p.rect.width * 25.4 / 72 for p in doc]
            scaled = min(sizes) * 170 / widths[0] if sizes else None
            assets.append(dict(path=pdf.relative_to(root).as_posix(), bytes=pdf.stat().st_size, sha256=sha(pdf), pages=len(doc), width_mm=widths[0], min_pdf_font_pt=min(sizes) if sizes else "raster_only", min_font_at_170mm_pt=scaled, embedded_fonts=embedded))
            check(pdf.stem + " PDF structure", len(doc) == 1 and bool(sizes) and embedded, "one page; selectable text; extractable non-Type3 fonts=" + str(embedded))
        png = pdf.with_suffix(".png")
        assets.append(dict(path=png.relative_to(root).as_posix(), bytes=png.stat().st_size, sha256=sha(png)))
    data = {}
    for n in range(1, 6):
        path = root / f"4_figures/Fig{n}/source_data/Figure_{n}_source_data.csv"
        data[n] = read_csv(path)
        assets.append(dict(path=path.relative_to(root).as_posix(), bytes=path.stat().st_size, sha256=sha(path)))
        check(f"F{n} source rows", bool(data[n]), len(data[n]))
    supplementary_sources = root / "phase17_v7/npj_sba_figure5_regulatory_ceiling/20260908_canonical_source_integration/figures/source_data"
    with zipfile.ZipFile(root / "6_release/SLE_Bcell_Submission_Package_2026-10-12/04_Supplementary_Data/Supplementary_Data_1_Figure_Source_Data.zip") as package:
        for n in range(1, 11):
            path = supplementary_sources / f"Supplementary_Figure_S{n}_source_data.csv"
            matches = [m for m in package.namelist() if m.endswith("/" + path.name) or m == path.name]
            if len(matches) != 1:
                raise ValueError("Supplementary Source Data archive membership not unique")
            check(f"S{n} archived-source identity", hashlib.sha256(package.read(matches[0])).hexdigest() == sha(path), matches[0])
            assets.append(dict(path=path.relative_to(root).as_posix(), bytes=path.stat().st_size, sha256=sha(path)))
    f3 = [r for r in data[3] if r["panel"] == "a"]
    check("Four-program family", len(f3) == 4, [r["program_id"] for r in f3])
    for row in f3:
        low, high = numbers["ci"][row["program_id"]]
        check("F3 CI algebra " + row["program_id"], abs(low - float(row["ci_low"])) < 1e-10 and abs(high - float(row["ci_high"])) < 1e-10, "saved effect/SE/df; no refit")
    control = [r for r in data[3] if r["panel"] == "d"]
    pan = next(r for r in control if r["analysis_name"] == "primary_base" and r["program_id"] == "PAN_B_IDENTITY_QC")
    check("Positive primary pan-B control retained", float(pan["ci_low"]) > 0 and pan["q_value_primary4"] == "", pan["effect"] + "; nominal-only QC")
    f4gate = [r for r in data[4] if r["panel"] == "d"]
    check("Calibration boundary", len(f4gate) == 3 and any(float(r["margin"]) < 0 for r in f4gate), [(r["metric"], r["observed"], r["criterion"]) for r in f4gate])
    genes = [r for r in data[4] if r["panel"] == "c"]
    rho = numbers["rho"]
    check("F4 shared genes and recomputed concordance", len(genes) == 4410 and round(float(rho), 3) == 0.026, f"n={len(genes)}; rho={rho:.9f}")
    joint = [r for r in genes if r["is_frozen_ifn_gene"].lower() == "true"]
    check("Ten jointly tested IFN genes", len(joint) == 10 and all(float(r["gse174188_logFC"]) > 0 and float(r["gse135779_logFC"]) > 0 for r in joint), len(joint))
    f5dep = [r for r in data[5] if r["panel"].lower() == "d"]
    check("Depletion records", len(f5dep) == 12, len(f5dep))
    exceptions = [r for r in f5dep if float(r["ci_low"]) < 0]
    check("Depletion negative interval not hidden", len(exceptions) == 1 and exceptions[0]["regulator"] == "STAT2", [(r["contrast"], r["regulator"], r["q_value"]) for r in exceptions])
    si = json.loads((root / "3_results/current/StageI_Result.json").read_text(encoding="utf-8"))
    wp = read_csv(root / "3_results/current/StageK_Source_Data.csv")
    check("Stage I same estimand as WP1", abs(float(wp[0]["point_estimate_T0"]) - si["point_estimate"]) < 1e-15, "not independent replication")
    check("Stage I Wald identities", abs(si["point_estimate"] / si["standard_error_HC3"] - si["test_statistic"]) < 1e-12 and abs(si["point_estimate"] - si["critical_value"] * si["standard_error_HC3"] - si["ci_lower"]) < 1e-12, "direct saved-value algebra; not refit")
    check("Stage I independent nominal P", abs(numbers["p"] - si["two_sided_p_value"]) < 1e-12, "SciPy t distribution; no model refit")
    save_csv(report / "SOURCE_AND_UNIT_AUDIT.csv", checks)
    save_json(report / "PROTECTED_ASSET_INVENTORY.json", assets)
    save_csv(report / "FIGURE_PDF_ARTWORK_METRICS.csv", [r for r in assets if r.get("pages")])
    # Preserve exact QA output, including failures, rather than turning the Skill's specification count into tests.
    skill = root / "8_reference/qiteng_v03242" / SKILL_NAME
    command = [sys.executable, str(skill / "04_EXECUTABLE_QA/qiteng_qa.py"), "package", str(skill)]
    run = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
    qa = json.loads(run.stdout)
    save_json(report / "QITENG_PACKAGE_QA_RAW.json", qa)
    mismatches = qa.get("manifest", {}).get("mismatches", [])
    zip_path = Path("C:/Users/Administrator/Desktop") / (SKILL_NAME + ".zip")
    with zipfile.ZipFile(zip_path) as z:
        prefix = SKILL_NAME + "/"
        archive_matches = []
        for name in ["00_CONTROL/VERSION_REGISTRY.json", "01_RUNTIME/CORE/SKILL.md", "04_EXECUTABLE_QA/qiteng_qa.py"]:
            raw = z.read(prefix + name)
            archive_matches.append(dict(path=name, zip_sha=hashlib.sha256(raw).hexdigest(), extracted_sha=sha(skill / name)))
    save_json(report / "QITENG_AUTHORITY_AND_LIMITATION.json", dict(zip_sha256=sha(zip_path), package_qa_exit=run.returncode, manifest_mismatches=len(mismatches), active_core_mismatches=[p for p in mismatches if p.startswith("01_RUNTIME/CORE/")], archive_matches=archive_matches, scientific_validation=False, behavioral_specs=qa.get("behavioral_spec_count")))
    print(json.dumps(dict(source_checks=len(checks), holds=sum(c["status"] != "PASS" for c in checks), protected_assets=len(assets), skill_manifest_mismatches=len(mismatches)), indent=2))


def render_pages(pdf, destination):
    destination.mkdir(parents=True, exist_ok=True)
    with fitz.open(pdf) as doc:
        for i, p in enumerate(doc):
            p.get_pixmap(dpi=120).save(destination / f"page-{i+1:02}.png")
        return len(doc)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", type=Path, default=Path.cwd())
    parser.add_argument("--render-only", action="store_true")
    args = parser.parse_args()
    root = args.workspace.resolve()
    out = root / "5_manuscript/scientific_narrative_candidate_2026-10-12"
    report = root / "7_report/scientific_narrative_reaudit_2026-10-12"
    report.mkdir(parents=True, exist_ok=True)
    out.mkdir(parents=True, exist_ok=True)
    if args.render_only:
        result = {name: render_pages(out / (name + ".pdf"), report / "render_qa" / name) for name in ["Manuscript", "Supplementary_Information"]}
        save_json(report / "RENDER_PAGE_COUNTS.json", result)
        print(result)
    else:
        audit(root, report)
        candidate = build(root, out, report)
        print("Candidate built; WPS rendering and visual QA remain required.")


if __name__ == "__main__":
    main()
