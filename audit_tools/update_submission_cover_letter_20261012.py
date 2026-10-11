#!/usr/bin/env python3
"""Create the title-synchronized 12 October 2026 cover-letter DOCX."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "04_submission/20260929_qiteng_submission_candidate/SLE_Bcell_Submission_Package/Cover_Letter/Cover_Letter.docx"
OUTPUT = ROOT / "01_manuscript/submission_current/Cover_Letter_npj_SBA_2026-10-12.docx"
OLD_TITLE = "Disease-blind reconstruction distinguishes reproducible interferon remodeling from less stable B-cell state assignments in systemic lupus erythematosus"
NEW_TITLE = "Interferon remodeling remains reproducible despite less stable B cell state assignments in systemic lupus erythematosus"


def replace_paragraph_text(paragraph, expected: str, replacement: str) -> None:
    if paragraph.text != expected:
        raise ValueError(f"fail-closed paragraph mismatch: {paragraph.text!r}")
    paragraph.clear()
    paragraph.add_run(replacement)


def main() -> None:
    document = Document(SOURCE)
    replace_paragraph_text(document.paragraphs[1], "29 September 2026", "12 October 2026")
    expected = f'Please consider our Article, "{OLD_TITLE}", for publication in npj Systems Biology and Applications.'
    replacement = f'Please consider our Article, "{NEW_TITLE}", for publication in npj Systems Biology and Applications.'
    replace_paragraph_text(document.paragraphs[4], expected, replacement)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    reread = Document(OUTPUT)
    if reread.paragraphs[1].text != "12 October 2026" or NEW_TITLE not in reread.paragraphs[4].text:
        raise ValueError("saved cover letter did not preserve the synchronized date/title")
    print(OUTPUT)


if __name__ == "__main__":
    main()
