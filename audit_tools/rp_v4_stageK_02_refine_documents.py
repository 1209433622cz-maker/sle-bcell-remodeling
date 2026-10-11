#!/usr/bin/env python3
"""Apply the independently reviewed Stage K wording to candidate DOCX files."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Pt


MAIN_OLD = {
    "In a later, outcome-exposed sensitivity analysis on the same 68 donors and the same previously observed contrast, fixed-reference profile mixing yielded a modestly larger managed-SLE association than pooled-count scoring (difference 0.034925 T0 units; HC3–t(64) 95% interval 0.011834–0.058016; Supplementary Table S10). This is a re-interval of an existing measurement-rule contrast, not an independent replication, a mediated fraction or a cell-intrinsic effect; the G2C simulation qualification is conditional on the tested generating mechanisms.":
    "In a later, outcome-exposed sensitivity analysis on the same 68 donors and the same previously observed contrast, fixed-reference profile mixing yielded a modestly larger managed-SLE association than pooled-count scoring (difference 0.034925 T0 units; HC3-t(64) 95% interval 0.011834-0.058016; Supplementary Table S10). This analysis re-estimated uncertainty for the existing measurement-rule contrast; it is not an independent replication, a mediated fraction or a cell-intrinsic effect. The G2C simulation qualification remains conditional on the tested generating mechanisms.",
    "A later, explicitly outcome-exposed measurement-rule sensitivity used a separate frozen 68-donor common-support subset (31 controls and 37 source-defined managed-SLE donors). For the same 12-gene T0 measurement, the disease-coefficient difference between fixed-reference linear CPM profile mixing and observed pooled-count mixing was re-intervalled with donor-level OLS, HC3 variance and t(64) critical values after synthetic operating-characteristic qualification. This is distinct from the 89-pseudobulk TMM-score primary model; exact definitions, dual intervals and analysis chronology are given in Supplementary Table S10.":
    "A later, explicitly outcome-exposed measurement-rule sensitivity used a separate frozen 68-donor common-support subset (31 controls and 37 source-defined managed-SLE donors). For the same 12-gene T0 measurement, we re-estimated uncertainty for the disease-coefficient difference between fixed-reference linear CPM profile mixing and observed pooled-count mixing using donor-level OLS, HC3 variance and t(64) critical values after synthetic operating-characteristic qualification. This is distinct from the 89-pseudobulk TMM-score primary model; exact definitions, both intervals and the analysis chronology are given in Supplementary Table S10.",
}

SUPP_OLD = {
    "Additional measurement-rule sensitivity—methods. In the frozen 68-donor common-support subset (31 controls and 37 source-defined managed-SLE donors), the original pooled-count score was calculated from summed raw IFN counts divided by each donor's full captured-UMI library total (CPM). For the paired fixed-profile score, five within-bin linear full-library CPM profiles were combined using the frozen normal-control pi_ref weights before the same 12-gene log2(CPM + 0.5) T0 transformation and standardization using the fixed 43-control reference means and population SDs. Their within-donor score difference was regressed on the same four original columns: intercept, managed-SLE indicator, age-centred covariate and frozen ethnicity indicator. HC3 variance and the two-sided t(64) critical value were used for the already observed coefficient difference. The G2C interval method was developed after the positive WP1 contrast was known, then evaluated in an independent synthetic validation (40 frozen scenarios: 30 applicable and 10 boundary reports) before the write-once real-data interval. The simulation result licenses this one interval only within the tested generating regimes and does not qualify other WP2 targets.":
    "Additional measurement-rule sensitivity: methods. In the frozen 68-donor common-support subset (31 controls and 37 source-defined managed-SLE donors), the original pooled-count score was calculated from summed raw IFN counts divided by each donor's full captured-UMI library total (CPM). For the paired fixed-profile score, five within-bin linear full-library CPM profiles were combined using the frozen normal-control pi_ref weights before the same 12-gene log2(CPM + 0.5) T0 transformation and standardization using the fixed 43-control reference means and population SDs. Their within-donor score difference was regressed on the same four original columns: intercept, managed-SLE indicator, age-centred covariate and frozen ethnicity indicator. HC3 variance and the two-sided t(64) critical value were used for the already observed coefficient difference. The G2C interval method was developed after the positive WP1 contrast was known, then evaluated in an independent synthetic validation (40 frozen scenarios: 30 applicable and 10 boundary reports) before the write-once real-data interval. The simulation result qualifies this one interval only within the tested generating regimes and does not qualify other WP2 targets.",
    "Additional measurement-rule sensitivity—results. On these same 68 donors the adjusted managed-SLE coefficient differed by 0.034925002 T0 units between fixed-profile and pooled-count scoring, equal within numerical precision to the previously observed RP v3 WP1 A1–A0 contrast. The earlier paired-percentile 95% bootstrap interval was 0.013789–0.057130 (2,000 paired donor resamples). The subsequent HC3–t(64) analysis of the identical contrast had a standard error of 0.011559, a 95% interval of 0.011834–0.058016, and a nominal two-sided P value of 0.003612. The procedures share the same donor outcomes, definition, reference and four-column design and must not be counted as independent confirmations.":
    "Additional measurement-rule sensitivity: results. On these same 68 donors the adjusted managed-SLE coefficient differed by 0.034925002 T0 units between fixed-profile and pooled-count scoring, equal within numerical precision to the previously observed RP v3 WP1 A1-A0 contrast. The earlier paired-percentile 95% bootstrap interval was 0.013789-0.057130 (2,000 paired donor resamples). The subsequent HC3-t(64) analysis of the identical contrast had a standard error of 0.011559, a 95% interval of 0.011834-0.058016, and a nominal two-sided P value of 0.003612. The procedures share the same donor outcomes, definition, reference and four-column design and must not be counted as independent confirmations.",
}


def replace_paragraphs(document: Document, replacements: dict[str, str], font_size: float) -> None:
    observed = set()
    for paragraph in document.paragraphs:
        if paragraph.text in replacements:
            observed.add(paragraph.text)
            paragraph.text = replacements[paragraph.text]
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
    missing = set(replacements) - observed
    if missing:
        raise ValueError(f"missing expected DOCX paragraphs: {len(missing)}")


def run(candidate: Path, output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    manuscript = Document(candidate / "Manuscript_StageK_Limited_Integration.docx")
    replace_paragraphs(manuscript, MAIN_OLD, 9.0)
    manuscript.save(output / "Manuscript_StageK_Limited_Integration.docx")

    supplement = Document(candidate / "Supplementary_Information_StageK_S10.docx")
    replace_paragraphs(supplement, SUPP_OLD, 8.4)
    replacements = {
        "Re-interval of same contrast; nominal P = 0.003612":
        "Uncertainty re-estimation for the same contrast; nominal P = 0.003612"
    }
    observed = 0
    for table in supplement.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text in replacements:
                        paragraph.text = replacements[paragraph.text]
                        for run in paragraph.runs:
                            run.font.size = Pt(7.2)
                        observed += 1
    if observed != 1:
        raise ValueError(f"expected one Stage I table-cell replacement, observed {observed}")
    supplement.save(output / "Supplementary_Information_StageK_S10.docx")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--candidate", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    run(args.candidate, args.output)
